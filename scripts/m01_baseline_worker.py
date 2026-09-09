"""Execute original v2 in its original interpreter; output only synthetic M01 evidence."""
import argparse
import ast
import dataclasses
import importlib.metadata
import json
import os
import pickle
import sqlite3
import subprocess
import sys
from pathlib import Path

from baseline_support import sha, write_json
from m01_baseline_support import SETTINGS


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('request')
    request = json.loads(Path(parser.parse_args().request).read_text('utf-8'))
    source, folder = Path(request['source_root']).resolve(), Path(request['output_dir']).resolve()
    root = Path(__file__).resolve().parents[1]
    assert folder.is_relative_to(root / 'output/validation/m01') and not folder.is_relative_to(source)
    assert Path.cwd().resolve() == folder and sys.dont_write_bytecode
    sys.path.insert(0, str(source))
    import numpy as np
    import pandas as pd
    import phonetic_toolbox
    from phonetic_toolbox.models.config import AcousticConfig
    from phonetic_toolbox.services import acoustic_service as service_module
    from phonetic_toolbox.services.io import lip as lip_module
    assert Path(phonetic_toolbox.__file__).resolve().is_relative_to(source)

    def pack(value):
        a = np.asarray(value)
        if a.dtype.kind in 'fiu':
            codes = np.where(np.isnan(a),1,np.where(np.isposinf(a),2,np.where(np.isneginf(a),3,0)))
            return {'values':np.where(np.isfinite(a),a,None).tolist(), 'nonfinite':codes.tolist(),
                    'shape':list(a.shape), 'dtype':str(a.dtype)}
        return {'values':a.tolist(), 'shape':list(a.shape), 'dtype':str(a.dtype)}

    def make_lip(mode='metadata', *, bad_length=False, insufficient=False):
        # Exact binary fractions make the expected affine tracks independently calculable.
        times = [0.5,0.0,0.25,0.25,0.75,float('nan')]
        data = {'metadata':{'lip_manual_offset':0.125},
                'relative_times':[x+10 for x in times],
                'absolute_timestamps':[x+100 for x in times]}
        for i,key in enumerate(['area','outer_width','open','circularity'],1):
            data[key] = [3*i,1*i,2*i,99*i,4*i,5*i]
        data['area'][2] = float('nan')  # interpolation bridges this missing source sample
        if mode == 'metadata': data['metadata']['audio_first_frame_time'] = 100.0
        path = folder / (mode + ('-bad' if bad_length else '-short' if insufficient else '') + '.pkl')
        if mode == 'companion':
            with path.with_name(path.stem+'_timestamps.pkl').open('wb') as f:
                pickle.dump({'start_time':100.0},f,protocol=4)
        if bad_length: data['outer_width'] = [1.0]
        if insufficient: data['relative_times']=[0.0]; data['absolute_timestamps']=[100.0]
        with path.open('wb') as f: pickle.dump(data,f,protocol=4)
        return path

    def gui_controls():
        from PyQt6 import QtWidgets
        from phonetic_toolbox.gui.dialogs.settings_dialog import SettingsDialog
        from phonetic_toolbox.gui.dialogs.parameter_tools_dialog import ParameterSelectionDialog
        from phonetic_toolbox.gui.widgets import parameter_estimation_widget as widget_module
        from phonetic_toolbox.services.settings_service import SettingsService
        app = QtWidgets.QApplication([])
        audit=json.loads((root/'docs/modules/evidence/M01-parameter-settings.json').read_text('utf-8'))
        dialog = SettingsDialog()
        result = {}
        for row in audit['settings']:
            control=getattr(dialog,row['control']); key=row['key']
            if row['range'] is None:
                default=control.isChecked(); bounds=None; control.setChecked(SETTINGS[key]); changed=control.isChecked()
            else:
                default=control.value(); bounds=[control.minimum(),control.maximum()]
                control.setValue(bounds[0]); lower=control.value()
                control.setValue(bounds[1]); upper=control.value()
                assert [lower,upper]==bounds
                control.setValue(int(SETTINGS[key]) if isinstance(control,QtWidgets.QSpinBox) else SETTINGS[key]); changed=control.value()
            result[key]={'default':default,'range':bounds,'changed':changed}
        dialog.save_settings()
        saved=dataclasses.asdict(SettingsService().get_config_object())
        assert all(saved[k]==v['changed'] for k,v in result.items())
        selection=ParameterSelectionDialog(list(service_module.PARAMETER_MAPPING.items()),['pF0'])
        selection._select_all(); assert len(selection.selected_keys())==80
        selection._select_none(); assert selection.selected_keys()==[]
        # Exercise actual widget event handler; replace only modal completion/message to avoid user UI.
        widget=widget_module.ParameterEstimationWidget(); previous=list(widget._selected_param_keys)
        warnings=[]
        QtWidgets.QMessageBox.warning=lambda *a: warnings.append(str(a[-1]))
        selection.exec=lambda: 1
        widget_module.ParameterSelectionDialog=lambda *a: selection
        widget._open_parameter_selection()
        rejected=bool(warnings) and widget._selected_param_keys==previous
        selection._select_all(); selection.exec=lambda: 0
        widget._selected_param_keys=['pF0']; widget._open_parameter_selection()
        cancelled=widget._selected_param_keys==['pF0']
        widget._player.stop()
        return {'settings':result,'saved':{k:saved[k] for k in SETTINGS},
                'empty_selection_rejected':rejected,'cancel_preserves_selection':cancelled,
                'scope':'offscreen original Qt controls and handlers; no visual/OS interaction claim'}

    mode=request.get('mode','analysis')
    if mode=='controls':
        scientific=gui_controls()
    elif mode=='lip':
        targets=np.arange(9,dtype=float)/8
        scientific={'target_times':targets.tolist(),'modes':{}}
        for mode_name in ['metadata','companion','relative']:
            path=make_lip(mode_name)
            scientific['modes'][mode_name]={k:pack(v) for k,v in lip_module.read_lip_data(str(path),targets).items()}
        scientific['insufficient']=lip_module.read_lip_data(str(make_lip(insufficient=True)),targets)
        scientific['bad_length']={k:pack(v) for k,v in lip_module.read_lip_data(str(make_lip(bad_length=True)),targets).items()}
    else:
        config=AcousticConfig()
        config.reaper_bin_path=str(source/'phonetic_toolbox/core/acoustic/reaper.exe')
        selection=request.get('selection')
        config.selected_parameter_keys=list(service_module.PARAMETER_MAPPING) if selection=='all' else selection
        for key,value in request.get('config',{}).items(): setattr(config,key,value)
        wav=Path(request['input']); inputs=[wav]
        lip_path=tg_path=None
        if request.get('associations'):
            lip_path=make_lip(); inputs.append(lip_path)
            from phonetic_toolbox.services.io.textgrid import TextGrid,Tier,Interval,write_textgrid
            tg_path=folder/'synthetic.TextGrid'
            tg=TextGrid(0.0,0.8,[Tier('音节',0.0,0.8,[Interval(0.0,0.8,'测试')]),
                Tier('IPA',0.0,0.8,[Interval(0.0,0.2,'aː'),Interval(0.2,0.4,'iː'),Interval(0.4,0.8,'=literal' if request.get('formula') else '标签')])])
            write_textgrid(tg,tg_path); inputs.append(tg_path)
        input_hashes={p.name:sha(p) for p in inputs}
        calls,returns,processes,lines,c_calls={}, {}, [], set(), []
        fault=request.get('fault')
        if fault=='irapt_raises':
            from phonetic_toolbox.core.acoustic import jitter_shimmer
            def fail_irapt(*args,**kwargs): raise RuntimeError('M01 controlled IRAPT unavailability')
            jitter_shimmer.irapt=fail_irapt
        original_run=subprocess.run
        def run_native(argv,*args,**kwargs):
            command=list(argv)
            if fault=='reaper_invalid_argument' and Path(command[0]).name.lower()=='reaper.exe':
                command=[command[0],'--m01-invalid-option'] # real native exit 1, then original Python fallback
            kwargs.setdefault('timeout',45)
            result=original_run(command,*args,**kwargs)
            processes.append({'executable_name':Path(command[0]).name,'executable_sha256':sha(command[0]),
                              'returncode':result.returncode})
            return result
        subprocess.run=run_native
        observed={'compute_energy','compute_silence_mask','compute_praat_f0_track','compute_reaper_f0',
            'compute_praat_formants','compute_spectral_features_batch','compute_jitter_shimmer',
            'smooth_preserving_gaps','smooth_lip','_extract_f0_for_wm','irapt','run_python_impl'}
        def profile(frame,event,value):
            module=frame.f_globals.get('__name__',''); name=frame.f_code.co_name
            if not module.startswith('phonetic_toolbox'): return
            if event=='call' and name in observed:
                scalars={k:v for k,v in frame.f_locals.items() if type(v) in (int,float,bool,str) and 'path' not in k and k not in ('reaper_bin',)}
                scalars={k:('<path>' if isinstance(v,str) and os.path.isabs(v) else v) for k,v in scalars.items()}
                calls.setdefault(name,[]).append(scalars)
            if event=='c_call' and name=='_extract_f0_for_wm': c_calls.append(getattr(value,'__name__',''))
            if event=='return' and name in {'irapt','_extract_f0_for_wm','run_python_impl'} and value is not None:
                values=value[2] if name=='run_python_impl' else value[0]
                a=np.asarray(values,dtype=float)
                returns[name]={'finite_positive':int((np.isfinite(a)&(a>0)).sum()),'values':pack(a)}
        def trace(frame,event,arg):
            if frame.f_code.co_name=='analyze_file' and frame.f_globals.get('__name__')==service_module.__name__:
                if event=='line': lines.add(frame.f_lineno)
                return trace
            return None
        sys.setprofile(profile); sys.settrace(trace)
        service=service_module.AcousticAnalysisService()
        try: result=service.analyze_file(str(wav),config,str(lip_path) if lip_path else None,str(tg_path) if tg_path else None)
        finally: sys.setprofile(None); sys.settrace(None); subprocess.run=original_run
        frame=result.to_dataframe()
        cfg=dataclasses.asdict(config); cfg['reaper_bin_path']='<v2>/phonetic_toolbox/core/acoustic/reaper.exe'
        setting=request.get('setting')
        setting_lines=[]
        if setting:
            tree=ast.parse(Path(service_module.__file__).read_text('utf-8'))
            setting_lines=[n.lineno for n in ast.walk(tree) if isinstance(n,ast.Attribute)
                and isinstance(n.value,ast.Name) and n.value.id=='config' and n.attr==setting]
        scientific={'config':cfg,'columns':list(frame.columns),'time_axis':result.time_axis.tolist(),
            'tracks':{str(c):pack(frame[c].to_numpy()) for c in frame.columns},
            'sampling_rate_metadata':result.sampling_rate,'observed_calls':calls,'backend_returns':returns,
            'native_processes':processes,'fault_injection':fault,'praat_wm_fallback_observed':bool(set(c_calls)&{'to_pitch','to_pitch_ac'}),
            'setting_observed':bool(set(setting_lines)&lines) if setting else None,
            'executed_setting_lines':sorted(set(setting_lines)&lines),'input_sha256':input_hashes}
        if request.get('export'):
            destination=folder/'synthetic.xlsx'
            exported={}
            try:
                service.save_results(result,str(destination))
                display=frame.rename(columns=service_module.PARAMETER_MAPPING)
                xlsx=pd.read_excel(destination)
                with sqlite3.connect(destination.with_suffix('.ptb.sqlite')) as db:
                    sqlite=pd.read_sql_query('SELECT * FROM params ORDER BY Time_s',db)
                    indexes=[row[1] for row in db.execute('PRAGMA index_list(params)')]
                def check(actual):
                    assert list(actual.columns)==list(display.columns)
                    for c in display:
                        if pd.api.types.is_numeric_dtype(display[c]):
                            np.testing.assert_array_equal(np.isnan(actual[c].to_numpy(dtype=float)),np.isnan(display[c].to_numpy(dtype=float)))
                            np.testing.assert_allclose(actual[c].to_numpy(dtype=float),display[c].to_numpy(dtype=float),rtol=1e-12,atol=1e-12,equal_nan=True)
                        else: assert actual[c].fillna('').tolist()==display[c].fillna('').tolist()
                    return True
                def compared(actual):
                    try: return check(actual)
                    except AssertionError: return False
                from openpyxl import load_workbook
                workbook=load_workbook(destination,data_only=False,read_only=True)
                formula_cells=sum(cell.data_type=='f' for row in workbook.active.iter_rows() for cell in row)
                workbook.close()
                exported={'status':'returned','xlsx_equal':compared(xlsx),'sqlite_equal':compared(sqlite),
                    'formula_cells':formula_cells,
                    'display_columns':list(display.columns),'sqlite_table':'params','sqlite_indexes':indexes}
            except Exception as exc:
                exported={'status':'error','error_type':type(exc).__name__,'xlsx_exists':destination.exists()}
            scientific['export']=exported
        assert all(sha(p)==input_hashes[p.name] for p in inputs)
    modules={}
    for name,module in list(sys.modules.items()):
        file=getattr(module,'__file__',None)
        if name.startswith('phonetic_toolbox') and file:
            path=Path(file).resolve(); assert path.is_relative_to(source)
            modules[name]=sha(path)
    dependencies={d.metadata['Name']:d.version for d in importlib.metadata.distributions()}
    output={'schema_version':1,'case_id':request['id'],
        'producer':{'root_kind':'original_v2','python':sys.version,'modules_sha256':modules,'installed_dependencies':dependencies,
            'recorder_sha256':{name:sha(root/'scripts'/name) for name in ['m01_baseline_worker.py','m01_baseline_support.py','baseline_support.py']}},
        'scientific':scientific}
    write_json(folder/'result.json.gz',output)
    print(json.dumps({'case':request['id'],'captured':True}),flush=True)


if __name__=='__main__': main()
