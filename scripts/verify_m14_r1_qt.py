"""M14-R1 hidden actual Qt, QWebChannel, owned synthetic tables and saved artifacts."""
import base64
import hashlib
import json
import os
import time
from pathlib import Path
os.environ.setdefault('QT_QPA_PLATFORM','windows')
os.environ.setdefault('QTWEBENGINE_CHROMIUM_FLAGS','--disable-gpu --disable-gpu-compositing --mute-audio')
from PyQt6.QtCore import QEventLoop,QTimer,Qt,QPoint
from PyQt6.QtTest import QTest
from PyQt6.QtWidgets import QApplication,QFileDialog
from ptb_desktop.host import Workbench,register_scheme
from verify_m14_wiring import setup,ROOT
from m14_r1_inputs import create

def main():
    out,db,cache=setup();print(out,flush=True);inputs=out/'inputs';create(inputs)
    saved=out/'saved';saved.mkdir();downloads=out/'downloads';downloads.mkdir()
    before={p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in inputs.iterdir()}
    register_scheme();app=QApplication(['M14-R1-owned-QA'])
    w=Workbench(ROOT/'frontend/dist',test=True,jobs_path=db,local_files_root=cache,vocal_profile=out/'vocal-profile')
    w.resize(1920,1080);w.setAttribute(Qt.WidgetAttribute.WA_DontShowOnScreen);w.show();w.page.setAudioMuted(True)
    QFileDialog.getExistingDirectory=lambda *a,**k:str(saved)
    QFileDialog.getSaveFileName=lambda *a,**k:(str(downloads/Path(a[2]).name),'')
    report=dict(success=False,checks=[],layouts=[],scope='Windows actual hidden Qt/QWebChannel, built frontend, real worker/files; synthetic only; no DDL')
    def pause(ms=100):
        loop=QEventLoop();QTimer.singleShot(ms,loop.quit);loop.exec()
    def js(code):
        loop=QEventLoop();box=[];w.page.runJavaScript(code,lambda v:(box.append(v),loop.quit()));QTimer.singleShot(6000,loop.quit);loop.exec()
        if not box:raise RuntimeError('JS timeout')
        return box[0]
    def until(code,seconds=90):
        end=time.monotonic()+seconds
        while time.monotonic()<end:
            if js(code):return
            pause()
        raise AssertionError(code+' '+str(js('document.body.innerText.slice(-4000)')))
    def click(text):
        assert js('(()=>{const e=[...document.querySelectorAll("button")].find(b=>b.offsetParent&&(b.getAttribute("aria-label")||b.textContent.trim())==='+json.dumps(text)+');if(!e||e.disabled)return false;e.click();return true})()'),text
        pause(120)
    def select(label,value):
        assert js('(()=>{const e=[...document.querySelectorAll("select")].find(e=>e.offsetParent&&(e.getAttribute("aria-label")==='+json.dumps(label)+'||e.closest("label")?.textContent.trim().startsWith('+json.dumps(label)+')));if(!e)return false;e.value='+json.dumps(str(value))+';e.dispatchEvent(new Event("change",{bubbles:true}));return true})()'),label
        pause()
    def fill(label,value,native=False):
        code='document.querySelector('+json.dumps('input[aria-label="'+label+'"]')+')'
        if native:
            pos=js('(()=>{const e='+code+';e.scrollIntoView({block:"nearest"});const r=e.getBoundingClientRect();e.focus();return {x:r.x+r.width/2,y:r.y+r.height/2}})()')
            target=w.view.focusProxy() or w.view;QTest.mouseClick(target,Qt.MouseButton.LeftButton,Qt.KeyboardModifier.NoModifier,QPoint(round(pos['x']),round(pos['y'])))
            QTest.keyClick(target,Qt.Key.Key_A,Qt.KeyboardModifier.ControlModifier);QTest.keyClicks(target,str(value));QTest.keyClick(target,Qt.Key.Key_Tab)
        else:
            js('(()=>{const e='+code+';e.value='+json.dumps(str(value))+';e.dispatchEvent(new Event("input",{bubbles:true}));e.dispatchEvent(new Event("change",{bubbles:true}));})()')
        pause()
        actual=js(code+'.value');assert actual==str(value),(label,value,actual)
    def idle():until('document.querySelector(".phonology-page")?.getAttribute("aria-busy")==="false"')
    def upload(name):
        raw=base64.b64encode((inputs/name).read_bytes()).decode()
        js('(()=>{const d=new DataTransfer();d.items.add(new File([Uint8Array.from(atob('+json.dumps(raw)+'),c=>c.charCodeAt(0))],'+json.dumps(name)+'));const e=document.querySelector(".phonology-page input[type=file]");e.files=d.files;e.dispatchEvent(new Event("change",{bubbles:true}));})()')
        idle();until('!!document.querySelector("table[aria-label=原始表格样例]")')
    def absent():assert not js('!!document.querySelector(".review")')
    def screenshot(name):
        w.view.repaint();pause(1000);w.view.grab();pause(500);app.processEvents();w.view.grab().save(str(out/name))
    try:
        until('!!document.querySelector("nav")');click('音系归纳');until('!!document.querySelector(".phonology-page")');absent()
        upload('table.docx');select('工作表 / 表格',1);idle();select('备注列',2);select('字头列',4);select('IPA 列',3);click('确认导入并继续');idle()
        until('document.body.innerText.includes("已读取 5 条记录")');absent()
        report['checks'].append('actual DOCX upload, second-table inspection and selected-column parsing')
        click('1 · 导入');select('文本编码','gb18030');select('分隔符','semicolon');upload('gbk.txt');select('字头列',1);select('IPA 列',2);select('备注列',3);click('确认导入并继续');idle();until('document.body.innerText.includes("已读取 2 条记录")')
        report['checks'].append('actual GB18030 semicolon text')
        click('1 · 导入');upload('large.xlsx');select('工作表 / 表格',1);idle();select('备注列',2);select('字头列',4);select('IPA 列',3);fill('数据开始行',3,True);idle();click('确认导入并继续');idle();until('document.body.innerText.includes("已读取 360 条记录")')
        fill('调类 35','yang',True);fill('调类 55','yin',True);click('上移调值 55');click('保存调类设置并继续');absent()
        assert not js('!!document.querySelector("dialog[open]")')
        js('[...document.querySelectorAll("[aria-label=声母列表] [role=option]")].find(e=>e.textContent==="ts").click()');click('归并声母');select('归并目标','m');click('确认归并');click('撤销归并 ts')
        js('[...document.querySelectorAll("[aria-label=声母列表] [role=option]")].find(e=>e.textContent==="ts").click()');click('归并声母');select('归并目标','m');click('确认归并')
        # Switching to import must preserve this uncommitted symbol edit.
        click('1 · 导入');click('3 · 声韵排序归并');assert not js('[...document.querySelectorAll("[aria-label=声母列表] [role=option]")].some(e=>e.textContent==="ts")')
        click('保存声韵设置并继续');until('!!document.querySelector(".result-sheet")')
        fill('搜索字头或 IPA','罕');until('document.querySelector(".result-sheet .preview-entry.focused")?.textContent.includes("罕")')
        fill('搜索字头或 IPA','罕 ');until('document.querySelector(".result-sheet .preview-entry.focused")?.textContent.includes("罕")')
        for label in ['韵母 → 声母DOCX','二维声韵表XLSX','声母 → 韵母DOCX']:
            click(label);until('document.querySelector(".result-sheet .preview-entry.focused")?.textContent.includes("罕")')
        fill('搜索字头或 IPA','ma',True);until('document.querySelector(".preview-caption").textContent.includes("命中")');click('下一个命中')
        fill('搜索字头或 IPA','');click('生成三份结果');idle();until('document.body.innerText.includes("三份结果已完整生成")')
        click('选择目录保存三份结果');idle();until('document.body.innerText.includes("三份结果已保存")');assert len(list(saved.iterdir()))==3
        hashes={p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in saved.iterdir()};click('选择目录保存三份结果');idle();until('document.body.innerText.includes("同名结果")');assert hashes=={p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in saved.iterdir()}
        report['checks'].append('native QTest number/tone/search keys, inline save/merge/undo and draft retained across steps; three preview modes, full-record search and real export/save/collision')
        # Download each exact saved file through real QWebEngine blob/save dialog adapter.
        for index in range(3):
            js('document.querySelectorAll(".result-files button")['+str(index)+'].click()');idle()
            end=time.monotonic()+20
            while time.monotonic()<end and len(list(downloads.iterdir()))<index+1:pause()
            assert len(list(downloads.iterdir()))==index+1
        end=time.monotonic()+20
        while time.monotonic()<end and any(not p.exists() or hashlib.sha256(p.read_bytes()).hexdigest()!=sha for name,sha in hashes.items() for p in [downloads/name]):pause()
        assert hashes=={p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in downloads.iterdir()}
        report['checks'].append('actual Qt individual DOCX/DOCX/XLSX blob downloads equal native directory save byte for byte')
        from docx import Document
        from openpyxl import load_workbook
        for path in saved.glob('*.docx'):
            text=''.join(p.text for p in Document(path).paragraphs);assert '罕' in text and text.count('注')==360
        sheet=load_workbook(next(saved.glob('*.xlsx')),rich_text=True).active
        assert sum(str(c.value).count('注') for row in sheet for c in row if c.value)==360
        report['checks'].append('all 360 notes and distant character read back from both DOCX files and XLSX')
        for width,height in [(1920,1080),(1366,768),(1100,800)]:
            w.resize(width,height);pause(200)
            for theme in ('light','dark'):
                js('document.documentElement.dataset.theme='+json.dumps(theme));pause(200)
                for s,name in enumerate(['导入','调类与调值','声韵排序归并','结果'],1):
                    click(str(s)+' · '+name)
                    if s<4:absent()
                    geometry=js('(()=>{const e=document.querySelector(".phonology-page");return {scroll:e.scrollWidth,width:e.clientWidth,dialogs:document.querySelectorAll("dialog[open]").length}})()')
                    assert geometry['scroll']<=geometry['width']+1 and geometry['dialogs']==0,geometry
                    screenshot(f'qt-r1-{width}-{theme}-step{s}.png');report['layouts'].append(dict(viewportWidth=width,viewportHeight=height,theme=theme,step=s,**geometry))
        w.resize(1920,1080);js('document.body.style.zoom="1.5"');pause(200)
        for s,name in enumerate(['导入','调类与调值','声韵排序归并','结果'],1):
            click(str(s)+' · '+name);g=js('(()=>{const e=document.querySelector(".phonology-page");return {scroll:e.scrollWidth,width:e.clientWidth}})()');assert g['scroll']<=g['width']+1,g
            screenshot(f'qt-r1-150pct-step{s}.png');report['layouts'].append(dict(zoom=1.5,step=s,**g))
        js('document.body.style.zoom="1"');pause(400);click('2 · 调类与调值');fill('调类 55','yin2',True);click('保存草稿');click('关闭 音系归纳');until('!document.querySelector(".phonology-page")');click('音系归纳');until('document.body.innerText.includes("已恢复本机草稿")');assert js('document.querySelector('+json.dumps('input[aria-label="调类 55"]')+').value')=='yin2'
        screenshot('qt-r1-reopened.png')
        assert before=={p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in inputs.iterdir()}
        report['checks'].append('24 actual Qt layouts plus four browser-zoom layouts, local draft reopen; source fixture hashes unchanged')
        report['success']=True
    finally:
        w.view.grab().save(str(out/'qt-r1-last.png'))
        (out/'qt-r1-report.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf8')
        js('window.onbeforeunload=null');w.close();pause(300);app.quit();print(json.dumps(report,ensure_ascii=False),flush=True)

if __name__=='__main__':main()
