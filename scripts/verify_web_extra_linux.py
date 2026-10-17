"""M07/M14 Linux admission: unchanged core, framed outputs and Windows parity."""
from pathlib import Path
import hashlib,io,json,sys,tempfile,zipfile
import numpy as np
from scipy.io import wavfile
from ptb_api.m07_models import M07Request
from ptb_api.font_models import FigureFontSnapshot
from ptb_worker.m07_executor import unpack as unpack_m07
from ptb_worker.m14_executor import unpack as unpack_m14
from ptb_worker.io.limits import Limits

ID='00000000-0000-0000-0000-000000000001'
def sha(value):return hashlib.sha256(value).hexdigest()
def sources():
    result={};rate=16000;t=np.arange(rate)/rate;rng=np.random.default_rng(20261017)
    for key,f0 in [('source',140),('target',200)]:
        y=sum(np.sin(2*np.pi*f0*k*t)/k for k in range(1,20))*.07+rng.normal(0,.0001,len(t));y[:800]=0;y[-800:]=0
        buffer=io.BytesIO();wavfile.write(buffer,rate,np.int16(y*32767));result[key]=buffer.getvalue()
    return result
def header(action,blobs,**kw):
    refs={key:dict(asset_id=ID,sha256=sha(value)) for key,value in blobs.items() if key in ('source','target')}
    request=M07Request(project_id=ID,idempotency_key='native-extra',action=action,**refs,**kw).model_dump()
    return dict(request=request,request_hash='a'*64,inputs=[dict(role=role,size_bytes=len(raw),sha256=sha(raw)) for role,raw in blobs.items()])
def canonical_document(raw):
    with zipfile.ZipFile(io.BytesIO(raw)) as archive:
        return {n:archive.read(n).decode('utf-8') for n in archive.namelist() if n=='word/document.xml' or n.startswith('xl/worksheets/')}
def main():
    windows=sys.platform=='win32';out=Path(sys.argv[1]);out.mkdir(parents=True,exist_ok=False);baseline=Path(sys.argv[2]);checks=[];observed={};documents={}
    if not windows:
        from ptb_worker.native.linux_runtime import load_profile,run
        profile,_=load_profile();profile_sha=sha(profile.read_bytes())
    else:profile_sha=None
    with tempfile.TemporaryDirectory(prefix='extra-owned-',dir=out) as temporary:
        folder=Path(temporary)
        def m07(h,b):
            if windows:
                from ptb_worker.m07_science import compute
                return dict(compute(h,b))
            request=folder/'m07.json';request.write_bytes(json.dumps(h).encode()+b'\n'+b''.join(b[item['role']] for item in h['inputs']))
            ev={};raw=run('m07',request,folder,Limits(output_bytes=64_000_000,process_bytes=1_000_000_000,timeout_seconds=120),evidence=ev);assert ev['cleaned'];return dict(unpack_m07(raw,h['request_hash'],h['request']))
        blobs=sources();analysis=m07(header('analyze',blobs),blobs);meta=json.loads(analysis['m07.ptb.json']);checks.append('analyze')
        controls=meta['controls'];controls['source'][8]+=10
        blobs['analysis']=analysis['analysis.m07.json'];applied=m07(header('apply',blobs,analysis_job_id=ID,controls=controls),blobs);checks.append('apply')
        for name in ('source_f0','target_f0'):observed[name]=np.asarray(meta[name])
        blobs['analysis']=applied['analysis.m07.json']
        for reverse in (False,True):
            for kind in (1,2,3):
                result=m07(header('generate',blobs,analysis_job_id=ID,continuum_type=kind,reverse_direction=reverse,generation={'step_count':3}),blobs)
                parts=[wavfile.read(io.BytesIO(result[f'step{i:02d}.wav']))[1] for i in range(1,4)];rate,combined=wavfile.read(io.BytesIO(result['combined_steps.wav']));assert rate==11025;np.testing.assert_array_equal(combined,np.concatenate(parts));assert np.max(np.abs(combined))>0
                observed[f'audio-{kind}-{reverse}']=combined;checks.append(f'generate-{kind}-{reverse}')
        if not windows:
            from ptb_worker.io.scratch import Scratch
            _,settings=load_profile()
            with Scratch(settings['cache'],400_000) as scratch:
                b={k:v for k,v in blobs.items() if k!='analysis'};h=header('analyze',b,analysis={'f0_backend':'reaper'})
                h.update(native_scratch=str(scratch.create(b'','.wav')),reaper_binary=settings['reaper_binary'])
                reaper=json.loads(m07(h,b)['m07.ptb.json'])
                values=np.asarray(reaper['source_f0']);values=values[np.isfinite(values)&(values>0)];assert len(values)>100 and 125<np.median(values)<155
                checks.append('native_reaper_analysis')
        m07_checks=list(checks);checks=[]
        def m14(raw,name,config):
            if windows:
                from ptb_worker.m14_jobs import execute
                return execute(raw,name,config)
            request=folder/'m14.json';request.write_bytes(json.dumps(dict(config=config,name=name,sha256=sha(raw))).encode()+b'\n'+raw)
            ev={};bundle=run('m14',request,folder,Limits(output_bytes=32_000_000,process_bytes=1_000_000_000,timeout_seconds=60),evidence=ev);assert ev['cleaned'];return dict(unpack_m14(bundle,sha(raw),config['action']))
        from openpyxl import Workbook
        from docx import Document
        rows=[['字头','音标','备注'],['巴','pa55',''],['爬','pʰa35',''],['妈','ma55',''],['怕','pʰa51','']]
        raw=('\n'.join(','.join(row) for row in rows)).encode();config=dict(action='preview',skip_first_row=True,consonant_only_as_zero_initial=True,computation_revision='m14/2',character_column=1,ipa_column=2,note_column=3,start_row=2,table_index=0,encoding='auto',delimiter='auto')
        preview=json.loads(m14(raw,'owned.csv',config)['m14-preview.json']);documents['analysis']=preview;checks.append('csv_preview')
        book=Workbook();sheet=book.active
        for row in rows:sheet.append(row)
        buffer=io.BytesIO();book.save(buffer);book.close();assert json.loads(m14(buffer.getvalue(),'owned.xlsx',{**config,'action':'inspect'})['m14-preview.json']);checks.append('xlsx_inspect')
        doc=Document();table=doc.add_table(rows=0,cols=3)
        for row in rows:
            for cell,text in zip(table.add_row().cells,row):cell.text=text
        buffer=io.BytesIO();doc.save(buffer);assert json.loads(m14(buffer.getvalue(),'owned.docx',{**config,'action':'inspect'})['m14-preview.json']);checks.append('docx_inspect')
        outputs=m14(raw,'owned.csv',{**config,'action':'export','settings':preview['config'],'font':FigureFontSnapshot(fallback_policy='portable').model_dump()});assert len(outputs)==3
        for name,value in outputs.items():documents[name]=canonical_document(value)
        checks.append('three_exports_readback')
    if windows:
        np.savez_compressed(baseline,**observed);baseline.with_suffix('.json').write_text(json.dumps(documents,ensure_ascii=False),encoding='utf-8')
    else:
        expected=np.load(baseline,allow_pickle=False)
        for key,value in observed.items():np.testing.assert_allclose(value,expected[key],rtol=0,atol=1 if key.startswith('audio-') else 1e-5,equal_nan=True)
        assert documents==json.loads(baseline.with_suffix('.json').read_text('utf-8'))
    for phase,phase_checks in [('m07',m07_checks),('m14',checks)]:
        (out/(phase+'.json')).write_text(json.dumps(dict(success=True,phase=phase,profile_sha256=profile_sha,checks=phase_checks),indent=2));print(phase+' PASS',flush=True)

if __name__=='__main__':main()
