"""Capture M14 operation states in an owned maximized Windows Qt workbench.

Uses a read-only backup of the existing task schema, a public fixture copy and
new synthetic tables. Never changes source fixtures, product code or registry.
Registration of successful captures belongs to the main manual author.
"""
from __future__ import annotations

import base64
import hashlib
import json
import os
from pathlib import Path
import shutil
import sys
import time
import traceback
from uuid import uuid4

ROOT = Path(__file__).resolve().parents[2]
DIRS = ('desktop/src', 'backend/src', 'packages/phonetic_core/src', 'scripts')
for directory in DIRS:
    sys.path.insert(0, str(ROOT / directory))
os.environ['PYTHONPATH'] = os.pathsep.join(str(ROOT / p) for p in DIRS)
os.environ.setdefault('QT_QPA_PLATFORM', 'windows')
os.environ.setdefault('QTWEBENGINE_CHROMIUM_FLAGS', '--disable-gpu --disable-gpu-compositing --mute-audio')


def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def main() -> int:
    sys.stdout.reconfigure(encoding='utf-8')
    from PyQt6.QtCore import QEventLoop, QTimer, Qt
    from PyQt6.QtWidgets import QApplication, QFileDialog
    from openpyxl import Workbook, load_workbook
    from docx import Document
    from ptb_desktop.host import Workbench, register_scheme
    from verify_m14_wiring import setup

    out = ROOT / 'output/manual-work/m14-captures' / uuid4().hex
    out.mkdir(parents=True)
    inputs, captures, saved, retry = [out / p for p in ('inputs', 'screenshots', 'saved', 'retry')]
    for directory in (inputs, captures, saved, retry):
        directory.mkdir()
    fixture = ROOT / 'tests/fixtures/m14/public.xlsx'
    originals = {str(fixture): sha(fixture)}
    shutil.copyfile(fixture, inputs / 'public-copy.xlsx')
    assert sha(inputs / 'public-copy.xlsx') == originals[str(fixture)]
    records = [
        ('示例01', 'pha55', '送气记法演示', '巴'), ('示例02', 'pʰa35', '送气记法演示', '怕'),
        ('示例03', 'ma55', '同声韵不同调值', '妈'), ('示例04', 'ma35', '同声韵不同调值', '麻'),
        ('示例05', 'na214', '调类编辑演示', '拿'), ('示例06', 'na21', '调类编辑演示', '那'),
        ('示例07', 'la55', '列表排序演示', '拉'), ('示例08', 'la35', '列表排序演示', '啦'),
        ('示例09', 'tsa55', '连音符记法演示', '资'), ('示例10', 't͡sa35', '连音符记法演示', '滋'),
        ('示例11', 'ku55', '韵母排列演示', '姑'), ('示例12', 'ku35', '韵母排列演示', '古'),
        ('示例13', 'ŋu55', '零声母对照演示', '午'), ('示例14', 'ŋu214', '零声母对照演示', '五'),
        ('示例15', 'i55', '零声母演示', '衣'), ('示例16', 'm55', '单辅音策略演示', '姆'),
    ]
    wb = Workbook()
    wb.active.title = '使用说明'
    wb.active.append(['合成操作字表，不是方言调查结果'])
    ws = wb.create_sheet('演示调查字表')
    ws.append(['本表仅演示列映射与分类操作', '', '', ''])
    ws.append(['编号', 'IPA', '备注', '字头'])
    for record in records:
        ws.append(record)
    table = inputs / '演示列映射.xlsx'
    wb.save(table)
    text_file = inputs / '演示GBK.txt'
    text_file.write_bytes('字头;IPA;备注\n妈;ma55;分号与编码演示\n麻;ma35;分号与编码演示\n'.encode('gb18030'))
    input_hashes = {str(p): sha(p) for p in inputs.iterdir()}
    state_out, db, cache = setup()
    register_scheme()
    app = QApplication(['PTB-owned-M14-manual-capture'])
    window = Workbench(ROOT / 'frontend/dist', test=True, jobs_path=db,
                       local_files_root=cache, vocal_profile=out / 'vocal-profile', start_module='M14')
    window.setAttribute(Qt.WidgetAttribute.WA_DontShowOnScreen)
    window.showMaximized()
    window.page.setAudioMuted(True)
    destination = {'path': saved}
    QFileDialog.getExistingDirectory = lambda *a, **k: str(destination['path'])
    report = {'success': False, 'chapterId': 'm14', 'out': str(out),
              'scope': 'Actual maximized hidden Windows Qt, production QWebChannel/local worker and synthetic single-syllable tables; no natural accuracy, Office rendering, physical compositor/DPI, IME or Linux GUI claim',
              'captures': [], 'checks': [], 'errors': [], 'terminations': [],
              'ownedTaskState': str(state_out), 'originalSourceSHA256': originals,
              'inputs': input_hashes, 'exports': [],
              'sourceSnapshot': [{'path': str(p.relative_to(ROOT)), 'sha256': sha(p)}
                                 for p in sorted((ROOT / 'frontend/dist').rglob('*'))
                                 if p.is_file() and p.suffix in ('.html', '.js', '.css')]}
    window.page.renderProcessTerminated.connect(lambda status, code: report['terminations'].append([status.name, code]))

    def pause(ms=100):
        loop = QEventLoop()
        QTimer.singleShot(ms, loop.quit)
        loop.exec()

    def js(code):
        loop, box = QEventLoop(), []
        window.page.runJavaScript(code, lambda value: (box.append(value), loop.quit()))
        QTimer.singleShot(6000, loop.quit)
        loop.exec()
        if not box:
            raise RuntimeError('JavaScript timeout')
        return box[0]

    def until(code, seconds=60):
        deadline = time.monotonic() + seconds
        while time.monotonic() < deadline:
            if js(code):
                return
            pause()
        raise AssertionError(code + '\n' + str(js('document.body.innerText.slice(-4000)')))

    def click(label):
        expr = '[...document.querySelectorAll("button")].find(b=>b.offsetParent&&(b.getAttribute("aria-label")||b.textContent.trim())===' + json.dumps(label) + ')'
        until('!!(' + expr + ')&&!(' + expr + ').disabled')
        js(expr + '.click()')
        pause(120)

    def select(label, value):
        assert js('(()=>{const e=[...document.querySelectorAll("select")].find(e=>e.offsetParent&&(e.getAttribute("aria-label")===' + json.dumps(label) + '||e.closest("label")?.textContent.trim().startsWith(' + json.dumps(label) + ')));if(!e)return false;e.value=' + json.dumps(str(value)) + ';e.dispatchEvent(new Event("change",{bubbles:true}));return true})()'), label
        pause(150)

    def fill(label, value):
        assert js('(()=>{const e=document.querySelector(' + json.dumps('input[aria-label="' + label + '"]') + ');if(!e)return false;e.value=' + json.dumps(str(value)) + ';e.dispatchEvent(new Event("input",{bubbles:true}));e.dispatchEvent(new Event("change",{bubbles:true}));return true})()'), label
        pause(150)

    def idle():
        until('document.querySelector(".phonology-page")?.getAttribute("aria-busy")==="false"')

    def upload(path):
        raw = base64.b64encode(path.read_bytes()).decode()
        js('(()=>{const d=new DataTransfer();d.items.add(new File([Uint8Array.from(atob(' + json.dumps(raw) + '),c=>c.charCodeAt(0))],' + json.dumps(path.name) + '));const e=document.querySelector(".phonology-page input[type=file]");e.files=d.files;e.dispatchEvent(new Event("change",{bubbles:true}));})()')
        idle()
        until('!!document.querySelector("table[aria-label=原始表格样例]")')

    def symbol(value, ctrl=False):
        assert js('(()=>{const e=[...document.querySelectorAll("[aria-label=声母列表] [role=option]")].find(e=>e.textContent===' + json.dumps(value) + ');if(!e)return false;e.dispatchEvent(new MouseEvent("click",{bubbles:true,ctrlKey:' + str(ctrl).lower() + '}));return true})()'), value
        pause(100)

    def snapshot(name, caption, state, after):
        assert window.isMaximized(), 'Capture requires a maximized window'
        pause(400)
        window.view.repaint()
        window.view.grab()
        pause(400)
        app.processEvents()
        pix = window.view.grab()
        assert [pix.width(), pix.height()] == [2160, 1350], (pix.width(), pix.height())
        file = captures / (name + '.png')
        assert pix.save(str(file)), file
        report['captures'].append({'id': name, 'chapterId': 'm14', 'file': str(file),
                                   'caption': caption, 'state': state, 'placeAfter': after,
                                   'distribution': 'public', 'maximized': True,
                                   'window': [window.width(), window.height()],
                                   'frame': [window.frameGeometry().width(), window.frameGeometry().height()],
                                   'image': [pix.width(), pix.height()], 'devicePixelRatio': pix.devicePixelRatio(),
                                   'theme': js('document.documentElement.dataset.theme'),
                                   'palette': js('document.documentElement.dataset.palette'),
                                   'fontSize': js('getComputedStyle(document.documentElement).fontSize'),
                                   'sha256': sha(file)})
        print('Captured ' + name, flush=True)

    try:
        until('!!document.querySelector("nav")&&document.fonts.status==="loaded"')
        click('设置')
        click('浅色')
        until('document.documentElement.dataset.theme==="light"')
        click('音系归纳')
        until('!!document.querySelector(".phonology-page")')
        upload(table)
        select('工作表 / 表格', 1)
        idle()
        select('字头列', 4)
        select('IPA 列', 2)
        select('备注列', 3)
        fill('数据开始行', 3)
        idle()
        until('!document.querySelector(".import-card [role=alert]")')
        snapshot('m14-selected-columns-r5-light', '另有说明工作表的演示字表：选择第二个工作表，从第 3 行读取，并把第 4、2、3 列分别设为字头、IPA、备注。', 'actual inspected XLSX; staged, not yet formally imported', 'm14-advanced-import')
        select('字头列', 2)
        until('document.querySelector(".import-card [role=alert]")?.textContent.includes("不同")')
        assert js('[...document.querySelectorAll("button")].find(b=>b.textContent.trim()==="确认导入并继续").disabled')
        snapshot('m14-column-conflict-r5-light', '字头列与 IPA 列重叠时，页面给出不同列提示并禁用确认导入；修改列映射即可继续。', 'actual local validation failure; no replacement of confirmed source', 'm14-section-02')
        select('字头列', 4)
        click('确认导入并继续')
        idle()
        until('document.body.innerText.includes("已读取 16 条记录")')
        fill('调类 55', '演示A')
        fill('调类 35', '演示A')
        fill('调类 214', '演示B')
        fill('调类 21', '演示B')
        click('上移调值 55')
        click('上移调值 55')
        snapshot('m14-tone-edited-r5-light', '演示性调类编辑：55 与 35 填入同名演示A，214 与 21 填入演示B，并把 55 上移；保存前显示有未保存修改。', 'actual toneDraft edits; arbitrary demonstration names, no dialect analysis', 'm14-advanced-tone')
        click('保存调类设置并继续')
        symbol('ts')
        symbol('t͡s', True)
        until('document.querySelector(".symbol-panel small")?.textContent.includes("已选 2")')
        snapshot('m14-symbol-multiselect-r5-light', '按 Ctrl 选中两个声母候选，列表显示已选 2；上移、下移和拖动只调整展示顺序。', 'actual Ctrl multi-selection via UI event; no merge yet', 'm14-advanced-merge')
        symbol('ts')
        click('归并声母')
        select('归并目标', 't͡s')
        until('document.querySelector("[aria-label=归并确认]")?.textContent.includes("影响 1 条记录")')
        snapshot('m14-merge-confirm-r5-light', '归并确认区显示 ts 指向 t͡s、影响 1 条记录与例字资；核对目标后再确认，也可取消本次归并。', 'actual pending merge target and affected-record computation; original IPA unchanged', 'm14-advanced-merge')
        click('确认归并')
        until('[...document.querySelectorAll("button")].some(b=>b.getAttribute("aria-label")==="撤销归并 ts")')
        snapshot('m14-merge-applied-r5-light', '确认后 ts 退出活动声母列表，下方出现 ts → t͡s 及撤销按钮；仍须保存声韵设置才进入结果。', 'actual in-page merge applied to editor draft; not persisted until explicit save', 'm14-advanced-merge')
        click('保存声韵设置并继续')
        click('生成三份结果')
        idle()
        until('document.body.innerText.includes("三份结果已完整生成")')
        click('声母 → 韵母DOCX')
        js('[...document.querySelectorAll("details")].filter(e=>e.querySelector("summary")?.textContent.includes("统计")).forEach(e=>e.open=true)')
        pause(200)
        snapshot('m14-initial-preview-r5-light', '生成后的声母到韵母快照：展开类别统计，右栏列出三份已完整生成的文件。演示调类仅说明同名归组。', 'real generated snapshot, initial-to-final preview; screen layout is not DOCX paper pagination', 'm14-output-files')
        click('韵母 → 声母DOCX')
        snapshot('m14-final-preview-r5-light', '同一次生成的韵母到声母快照：以韵母为外层，依次展示声母及其演示调类下的字项。', 'real generated snapshot, final-to-initial preview; same 16 records and confirmed merge', 'm14-output-files')
        click('二维声韵表XLSX')
        snapshot('m14-matrix-preview-r5-light', '同一次生成的二维声韵表：行是韵母，列是声母，格内保留演示调类、字头及备注；空格仅反映本字表未收录组合。', 'real generated snapshot, matrix preview; not a claim of spreadsheet application rendering', 'm14-output-files')
        fill('搜索字头或 IPA', 'tsa55')
        js('[...document.querySelectorAll("details")].filter(e=>e.querySelector("summary")?.textContent.includes("导入与归并审阅")).forEach(e=>e.open=true)')
        until('document.querySelector(".review")?.textContent.includes("tsa55")')
        snapshot('m14-merged-review-r5-light', '搜索原始记音 tsa55 并展开审阅：来源行与原始 IPA 保留，声母列同时显示 ts 归并为 t͡s 的关系。', 'actual whole-record original IPA search with generated-snapshot review', 'm14-advanced-preview')
        fill('搜索字头或 IPA', '')
        click('选择目录保存三份结果')
        idle()
        until('document.body.innerText.includes("三份结果已保存")')
        assert len(list(saved.iterdir())) == 3
        hashes = {p.name: sha(p) for p in saved.iterdir()}
        for p in saved.iterdir():
            if p.suffix == '.docx':
                content = '\n'.join(paragraph.text for paragraph in Document(p).paragraphs)
                assert all(record[3] in content for record in records), p
                assert all(record[2] in content for record in records), p
                assert '演示A' in content and '演示B' in content
            elif p.suffix == '.xlsx':
                sheet = load_workbook(p, rich_text=True).active
                content = '\n'.join(str(c.value) for row in sheet for c in row if c.value is not None)
                assert sheet.title == '二维同音字表' and sheet.freeze_panes == 'B2'
                assert all(record[3] in content for record in records), p
                assert all(record[2] in content for record in records), p
                assert '演示A' in content and '演示B' in content
            report['exports'].append({'file': str(p), 'sha256': sha(p), 'size': p.stat().st_size, 'recordsReadBack': 16})
        report['checks'].append('All 16 synthetic characters and demonstration classes/notes read back from DOCX/DOCX/XLSX; XLSX sheet/freeze verified')
        click('选择目录保存三份结果')
        idle()
        until('document.body.innerText.includes("同名结果")')
        assert hashes == {p.name: sha(p) for p in saved.iterdir()}
        snapshot('m14-save-conflict-r5-light', '再次选到已有同名成果的目录时，状态栏拒绝覆盖；原三份文件保持不变，结果仍在页内可重试保存。', 'actual native directory save collision; prior output SHA unchanged', 'm14-section-02')
        destination['path'] = retry
        click('选择目录保存三份结果')
        idle()
        until('document.body.innerText.includes("三份结果已保存")')
        assert hashes == {p.name: sha(p) for p in retry.iterdir()}
        report['checks'].append('Real output-directory collision refusal and retry to a fresh directory; complete files equal byte for byte')
        click('1 · 导入')
        select('文本编码', 'gb18030')
        select('分隔符', 'semicolon')
        upload(text_file)
        select('字头列', 1)
        select('IPA 列', 2)
        select('备注列', 3)
        fill('数据开始行', 2)
        idle()
        snapshot('m14-gbk-delimiter-r5-light', 'GBK 文本的原始样例：编码选 GB18030 / GBK，分隔符选分号，检查三列内容后再确认导入。', 'actual GB18030 semicolon sample; staged alternate input leaves earlier confirmed result intact', 'm14-advanced-import')
        click('取消此次选择')
        click('4 · 结果')
        until('document.querySelectorAll(".result-files li").length===3')
        assert input_hashes == {str(p): sha(p) for p in inputs.iterdir()}
        assert originals == {str(fixture): sha(fixture)}
        assert not report['terminations']
        report['checks'].append('Original public fixture and every isolated input SHA preserved; cancelling pending alternate input retains confirmed generated result')
        report['success'] = True
    except Exception:
        report['errors'].append(traceback.format_exc())
    finally:
        (out / 'report.json').write_text(json.dumps(report, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
        window.close()
        pause(300)
        app.quit()
        print(json.dumps({'success': report['success'], 'report': str(out / 'report.json'), 'captures': len(report['captures']), 'errors': report['errors']}, ensure_ascii=False), flush=True)
    return 0 if report['success'] else 1


if __name__ == '__main__':
    raise SystemExit(main())
