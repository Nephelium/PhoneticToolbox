"""Count every source character, including duplicates, in three actual exports."""
from collections import Counter
import argparse
import json
from pathlib import Path
import sys


def main():
    sys.stdout.reconfigure(encoding='utf-8')
    parser = argparse.ArgumentParser()
    parser.add_argument('--capture-report', type=Path, required=True)
    args = parser.parse_args()
    from docx import Document
    from openpyxl import load_workbook
    report = json.loads(args.capture_report.read_text(encoding='utf-8'))
    assert report['success'] is True
    rows = list(load_workbook(report['sourceCopy'], data_only=True, read_only=True).active.values)
    expected = Counter(str(row[0]) for row in rows)
    checks = {}
    for item in report['exports']:
        path = Path(item['file'])
        actual, records = Counter(), 0
        if path.suffix == '.xlsx':
            sheet = load_workbook(path, rich_text=True).active
            for row in list(sheet)[1:]:
                for cell in row[1:]:
                    if cell.value is None:
                        continue
                    for block in cell.value:
                        if (hasattr(block, 'font') and not block.font.vertAlign
                                and not block.text.startswith('[') and block.text.strip()):
                            actual.update(block.text)
                            records += 1
        else:
            document = Document(path)
            for paragraph in document.paragraphs:
                if '[' not in paragraph.text:
                    continue
                for run in paragraph.runs:
                    if len(run.text) == 1 and run.text in expected and not run.font.subscript:
                        actual.update(run.text)
                        records += 1
        assert actual == expected and records == len(rows), (path.name, records, actual - expected, expected - actual)
        checks[path.name] = dict(records=records, characterCounterExact=True)
    target = args.capture_report.parent / 'record-count-readback.json'
    target.write_text(json.dumps(checks, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    print(json.dumps(checks, ensure_ascii=False))


if __name__ == '__main__':
    main()
