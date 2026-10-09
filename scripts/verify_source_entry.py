"""P19: real source workers and hidden Qt, with new owned validation outputs."""
import argparse
import hashlib
import io
import json
import os
from pathlib import Path
import struct
import subprocess
import sys
import time
import traceback

from workbench_source import ROOT, configure


def verify(out):
    out = Path(out).resolve()
    out.mkdir(parents=True, exist_ok=False)
    report = dict(success=False, binding=configure(), workers={}, modules=[])
    window = None
    try:
        import numpy as np
        from scipy.io import wavfile
        from ptb_worker.m14_jobs import execute
        from docx import Document
        from openpyxl import load_workbook
        table = (ROOT / 'tests/fixtures/m14/public.xlsx').read_bytes()
        settings = dict(action='preview', skip_first_row=True, consonant_only_as_zero_initial=True)
        preview = json.loads(execute(table, 'public.xlsx', settings)['m14-preview.json'])
        files = execute(table, 'public.xlsx', settings | dict(action='export', settings=preview['config'],
                        font=dict(schema_version='font/1', zh='Microsoft YaHei', latin='Segoe UI', ipa='Doulos SIL', size_px=14)))
        for name, raw in files.items():
            (out / name).write_bytes(raw)
            if name.endswith('.docx'):
                assert Document(io.BytesIO(raw)).tables
            else:
                book = load_workbook(io.BytesIO(raw), read_only=True)
                assert book.sheetnames
                book.close()
        report['document_exports'] = {name:len(raw) for name,raw in files.items()}
        def child(kind, python, entry, args):
            code = '''import json,runpy,sys
entry=sys.argv[1];sys.argv=sys.argv[1:];runpy.run_path(entry,run_name='__main__')
print(json.dumps({n:sys.modules[n].__file__ for n in ('phonetic_core','ptb_worker') if n in sys.modules}))
'''
            env = dict(os.environ, LOCALAPPDATA=str(out / 'private-localappdata'))
            result = subprocess.run([str(python), '-I', '-B', '-X', 'utf8', '-c', code, str(entry), *map(str,args)],
                                    cwd=out, env=env, capture_output=True, encoding='utf8', timeout=150,
                                    creationflags=getattr(subprocess,'CREATE_NO_WINDOW',0))
            (out / (kind+'-stderr.txt')).write_text(result.stderr,'utf8')
            assert result.returncode == 0, result.stderr
            origins = json.loads(result.stdout.strip().splitlines()[-1])
            assert Path(origins['phonetic_core']) == ROOT / 'packages/phonetic_core/src/phonetic_core/__init__.py'
            assert Path(origins['ptb_worker']) == ROOT / 'backend/src/ptb_worker/__init__.py'
            report['workers'][kind] = origins
        fixture = np.load(ROOT / 'tests/fixtures/m03/EGG-SYN-PCM16.npz')
        egg = io.BytesIO()
        wavfile.write(egg,44100,np.column_stack((fixture['load.audio_signal'],fixture['load.egg_signal_raw'])))
        tone = io.BytesIO()
        wavfile.write(tone,16000,(.25*np.sin(2*np.pi*150*np.arange(16000)/16000)).astype(np.float32))
        for kind,raw,settings in [('egg',egg.getvalue(),dict(mode='single',roi_start=0,roi_end=.5)),
                                 ('lpc',tone.getvalue(),dict(roi_start=.1,roi_end=.15,order=20))]:
            request, response = out / (kind+'.request'), out / (kind+'.response')
            header = dict(sha256=hashlib.sha256(raw).hexdigest(),config=settings,input_name='public.wav')
            if kind == 'lpc':
                header.update(audio_size=len(raw), textgrid_sha256=None)
            request.write_bytes(json.dumps(header).encode()+b'\n'+raw)
            child(kind,os.environ['PTB_EGG_PYTHON'],ROOT / f'backend/src/ptb_worker/{kind}_bootstrap.py',[request,response])
            data = response.read_bytes()
            length = struct.unpack('<Q',data[:8])[0]
            manifest = json.loads(data[8:8+length])
            assert not manifest.get('error'), manifest
            position = 8+length
            for file in manifest['files']:
                payload=data[position:position+file['size_bytes']]; position+=len(payload)
                assert hashlib.sha256(payload).hexdigest()==file['sha256']
            assert position==len(data)
            report['workers'][kind]['outputs']=manifest['files']
        import cv2
        media = out / 'media'
        media.mkdir()
        writer=cv2.VideoWriter(str(media/'public.mp4'),cv2.VideoWriter_fourcc(*'mp4v'),24,(128,128))
        assert writer.isOpened()
        try:
            for _ in range(12):writer.write(np.zeros((128,128,3),np.uint8))
        finally:writer.release()
        request=media/'request.json'
        request.write_text(json.dumps(dict(input='public.mp4',config=dict(filter_enabled=False,cutoff_hz=8,animation='none'))),'utf8')
        child('m05',os.environ['PTB_M05_PYTHON'],ROOT/'backend/src/ptb_worker/m05_child.py',[request])
        result=json.loads((media/'response.json').read_text('utf8'))
        assert result['success'] and result['frames']==12,result
        report['workers']['m05']['result']=result
        os.environ['QT_QPA_PLATFORM']='windows'
        os.environ['QTWEBENGINE_CHROMIUM_FLAGS']='--mute-audio'
        from PyQt6.QtCore import QEventLoop,QTimer,Qt,QUrl
        from PyQt6.QtWidgets import QApplication
        from ptb_desktop.host import Workbench,register_scheme
        register_scheme()
        app=QApplication(['source-entry-owned-QA'])
        app.setApplicationName('source-entry-owned-QA')
        window=Workbench(ROOT/'frontend/dist',test=True,vocal_profile=out/'vocal',start_module='M14')
        window.setAttribute(Qt.WidgetAttribute.WA_DontShowOnScreen)
        window.show()
        def pause():
            loop=QEventLoop();QTimer.singleShot(100,loop.quit);loop.exec()
        def js(code):
            loop=QEventLoop();values=[]
            window.page.runJavaScript(code,lambda v:(values.append(v),loop.quit()))
            QTimer.singleShot(5000,loop.quit);loop.exec()
            assert values,'JavaScript timeout'
            return values[0]
        for module in ('M14','M01','M05','M10','M16','M17'):
            # A fragment-only navigation does not mount a new Vue application.
            if module!='M14':window.view.load(QUrl('ptbapp://app/index.html?source-entry-case='+module+'#'+module))
            deadline=time.monotonic()+40
            while time.monotonic()<deadline:
                if js('document.querySelector("#tab-'+module+'")?.getAttribute("aria-selected")==="true"'):break
                pause()
            else:raise RuntimeError('Module did not open: '+module)
            report['modules'].append(module)
        window.vocal.start()
        report['vocal_worker_started']=window.vocal.process.poll() is None
        report['local_service_health']=window.service.get('/api/v1/health')
        service=window.service.process
        native=window.vocal.process
        window.closing=True;window.close();pause()
        assert service.poll()==0 and native.poll()==0
        report['owned_process_exit_codes']=[service.returncode,native.returncode]
        report['success']=True
    except Exception:
        report['error']=traceback.format_exc()
    finally:
        if window is not None:
            window.closing=True;window.close()
        (out/'report.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n','utf8')
    print(json.dumps({'success':report['success'],'output':str(out),'error':report.get('error')}),flush=True)
    return 0 if report['success'] else 1


if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output',type=Path,required=True)
    raise SystemExit(verify(parser.parse_args().output))
