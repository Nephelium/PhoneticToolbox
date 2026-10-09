"""Public synthetic M14-R1 files only, written beneath an owned validation root."""
import io
from pathlib import Path
import sys
from openpyxl import Workbook
from docx import Document

def create(root):
    root=Path(root);root.mkdir(parents=True,exist_ok=True)
    values=[['编号','备注','IPA','字头']]
    for i in range(360):
        values.append([str(i),f'注{i}','tsã35' if i==0 else 't͡s55' if i==1 else 'ma35' if i==2 else 'ma55','鼻' if i==0 else '罕' if i==319 else '妈'])
    w=Workbook();w.active.title='说明';w.active.append(['说明','不选此表']);s=w.create_sheet('调查字表');s.append(['调查说明']);
    for row in values:s.append(row)
    w.save(root/'large.xlsx')
    d=Document();d.add_table(rows=1,cols=1).cell(0,0).text='说明';t=d.add_table(rows=6,cols=4)
    for i,row in enumerate(values[:6]):
        for j,value in enumerate(row):t.cell(i,j).text=value
    d.save(root/'table.docx')
    (root/'gbk.txt').write_bytes('字;IPA;备注\n妈;ma55;中文备注\n麻;ma35;'.encode('gb18030'))
    (root/'missing.tsv').write_text('字\tIPA\t备注\n\tma55\t备注',encoding='utf8')

if __name__=='__main__':create(sys.argv[1])
