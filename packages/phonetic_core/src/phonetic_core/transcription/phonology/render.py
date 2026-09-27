# Direct V2 migration; source_id=PENDING-PHONOLOGY. See NOTICE.md and M14-source-map.md.
from __future__ import annotations
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from .models import ParsedPhonologyRow, PhonologyAnalysisResult
from .rules import PhonologyRules
_OPENPYXL_RICH_TEXT_AVAILABLE=True

class PhonologyRenderer(PhonologyRules):
    def _write_word_document(
        self,
        analysis: PhonologyAnalysisResult,
        tone_class_map: dict[str, str],
        output_path: BytesIO,
        mode: str,
        ordered_initials: list[str],
        ordered_finals: list[str],
        ordered_tones: list[str],
    ) -> None:
        Document = self._resolve_document_factory()

        doc = Document()
        self._setup_word_document_styles(doc)
        title = doc.add_heading("同音字表", level=1)
        self._center_word_paragraph(title)
        if mode == "final_initial":
            subtitle = doc.add_paragraph("（韵母 → 声母）")
            grouped = self._group_rows(
                analysis=analysis,
                tone_class_map=tone_class_map,
                outer_key="final",
                inner_key="initial",
            )
            outer_values = ordered_finals
            inner_values = ordered_initials
        else:
            subtitle = doc.add_paragraph("（声母 → 韵母）")
            grouped = self._group_rows(
                analysis=analysis,
                tone_class_map=tone_class_map,
                outer_key="initial",
                inner_key="final",
            )
            outer_values = ordered_initials
            inner_values = ordered_finals
        self._center_word_paragraph(subtitle)
        self._style_word_paragraph_runs(subtitle, size=11, bold=True)

        self._append_summary_sections(
            doc,
            analysis,
            tone_class_map,
            ordered_initials=ordered_initials,
            ordered_finals=ordered_finals,
            ordered_tones=ordered_tones,
        )
        for outer in outer_values:
            if outer not in grouped:
                continue
            category_label = outer if outer else "空韵"
            category = doc.add_heading(f"{category_label}", level=3)
            self._center_word_paragraph(category)
            self._style_word_paragraph_runs(category, size=12, bold=True)
            inner_map = grouped[outer]
            for inner in inner_values:
                if inner not in inner_map:
                    continue
                p = doc.add_paragraph()
                self._add_word_text(p, f"{inner} ", size=12, bold=True)
                tone_map = inner_map[inner]
                for tone_label in self._ordered_tone_labels(
                    set(tone_map.keys()), tone_class_map, ordered_tones
                ):
                    self._add_word_text(p, f"[{tone_label}]")
                    self._append_word_entries(p, tone_map[tone_label])
            spacer = doc.add_paragraph("")
            self._style_word_paragraph_runs(spacer)
        doc.save(output_path)

    def _resolve_document_factory(self):
        from docx import Document
        return Document

    def _append_summary_sections(
        self,
        doc,
        analysis: PhonologyAnalysisResult,
        tone_class_map: dict[str, str],
        ordered_initials: list[str],
        ordered_finals: list[str],
        ordered_tones: list[str],
    ) -> None:
        heading_initial = doc.add_heading("声母统计", level=2)
        self._center_word_paragraph(heading_initial)
        self._style_word_paragraph_runs(heading_initial, size=12, bold=True)
        initial_lines = []
        for initial in ordered_initials:
            examples = self._pick_examples(analysis.rows, lambda row: row.initial == initial)
            initial_lines.append((f"{initial}: ", examples))
        self._append_two_column_entries(doc, initial_lines)

        heading_final = doc.add_heading("韵母统计", level=2)
        self._center_word_paragraph(heading_final)
        self._style_word_paragraph_runs(heading_final, size=12, bold=True)
        final_lines = []
        for final in ordered_finals:
            label = final if final else "空韵"
            examples = self._pick_examples(analysis.rows, lambda row: row.final == final)
            final_lines.append((f"{label}: ", examples))
        self._append_two_column_entries(doc, final_lines)

        heading_tone = doc.add_heading("声调统计", level=2)
        self._center_word_paragraph(heading_tone)
        self._style_word_paragraph_runs(heading_tone, size=12, bold=True)
        tone_lines = []
        for tone_value in ordered_tones:
            tone_label = tone_class_map.get(tone_value, tone_value or "0")
            examples = self._pick_examples(
                analysis.rows, lambda row: row.tone_value == tone_value
            )
            tone_lines.append((f"{tone_value} → {tone_label}: ", examples))
        self._append_two_column_entries(doc, tone_lines)
        spacer = doc.add_paragraph("")
        self._style_word_paragraph_runs(spacer)

    def _pick_examples(
        self,
        rows: list[ParsedPhonologyRow],
        predicate,
        max_count: int = 5,
    ) -> list[ParsedPhonologyRow]:
        picked: list[ParsedPhonologyRow] = []
        seen: set[tuple[str, str]] = set()
        for row in rows:
            if not predicate(row):
                continue
            key = (row.character, row.note)
            if key in seen:
                continue
            seen.add(key)
            picked.append(row)
            if len(picked) >= max_count:
                break
        return picked

    def _append_word_entries(self, paragraph, entries: list[ParsedPhonologyRow]) -> None:
        for idx, row in enumerate(entries):
            if idx > 0:
                self._add_word_text(paragraph, " ")
            self._add_word_text(paragraph, row.character)
            if row.note:
                self._add_word_text(paragraph, row.note, size=12, subscript=True)

    def _write_matrix_xlsx(
        self,
        analysis: PhonologyAnalysisResult,
        tone_class_map: dict[str, str],
        output_path: BytesIO,
        ordered_initials: list[str],
        ordered_finals: list[str],
        ordered_tones: list[str],
    ) -> None:
        grouped = self._group_rows(
            analysis=analysis,
            tone_class_map=tone_class_map,
            outer_key="final",
            inner_key="initial",
        )
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "二维同音字表"
        header_fill = PatternFill("solid", fgColor="DCE6F1")
        thin_side = Side(style="thin", color="BFBFBF")
        thin_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
        center_alignment = Alignment(horizontal="center", vertical="center")
        body_alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)

        sheet.cell(row=1, column=1, value="韵母\\声母")
        sheet.cell(row=1, column=1).font = Font(name="Times New Roman", size=11, bold=True)
        sheet.cell(row=1, column=1).alignment = center_alignment
        sheet.cell(row=1, column=1).fill = header_fill
        sheet.cell(row=1, column=1).border = thin_border
        for col_idx, initial in enumerate(ordered_initials, start=2):
            sheet.cell(row=1, column=col_idx, value=initial)
            sheet.cell(row=1, column=col_idx).font = Font(name="Times New Roman", size=11, bold=True)
            sheet.cell(row=1, column=col_idx).alignment = center_alignment
            sheet.cell(row=1, column=col_idx).fill = header_fill
            sheet.cell(row=1, column=col_idx).border = thin_border

        for row_idx, final in enumerate(ordered_finals, start=2):
            sheet.cell(row=row_idx, column=1, value=final if final else "")
            sheet.cell(row=row_idx, column=1).font = Font(name="Times New Roman", size=11, bold=True)
            sheet.cell(row=row_idx, column=1).alignment = center_alignment
            sheet.cell(row=row_idx, column=1).fill = header_fill
            sheet.cell(row=row_idx, column=1).border = thin_border
            inner_map = grouped.get(final, {})
            for col_idx, initial in enumerate(ordered_initials, start=2):
                tone_map = inner_map.get(initial, {})
                cell = sheet.cell(row=row_idx, column=col_idx)
                cell.alignment = body_alignment
                cell.border = thin_border
                cell.font = Font(name="Times New Roman", size=11)
                if not tone_map:
                    continue
                rich_value = self._build_rich_text_cell_value(tone_map)
                if rich_value is not None:
                    cell.value = rich_value
                else:
                    chunks: list[str] = []
                    for tone_label in self._ordered_tone_labels(
                        set(tone_map.keys()), tone_class_map, ordered_tones
                    ):
                        chars = "".join(
                            self._render_plain_entry(entry) for entry in tone_map[tone_label]
                        )
                        chunks.append(f"[{tone_label}]{chars}")
                    cell.value = " ".join(chunks)
        for col_idx in range(1, len(ordered_initials) + 2):
            col_letter = get_column_letter(col_idx)
            sheet.column_dimensions[col_letter].width = 18 if col_idx == 1 else 26
        for row_idx in range(1, len(ordered_finals) + 2):
            sheet.row_dimensions[row_idx].height = 36 if row_idx == 1 else 52
        sheet.freeze_panes = "B2"
        workbook.save(output_path)

    def _ordered_tone_labels(
        self,
        labels: set[str],
        tone_class_map: dict[str, str],
        ordered_tones: list[str],
    ) -> list[str]:
        ordered_labels: list[str] = []
        seen: set[str] = set()
        for tone_value in ordered_tones:
            label = tone_class_map.get(tone_value, tone_value or "0")
            if label in labels and label not in seen:
                ordered_labels.append(label)
                seen.add(label)
        for label in self._sort_tones(labels):
            if label not in seen:
                ordered_labels.append(label)
                seen.add(label)
        return ordered_labels

    def _append_two_column_entries(self, doc, entries: list[tuple[str, list[ParsedPhonologyRow]]]):
        if not hasattr(doc, "add_table"):
            for label, examples in entries:
                p = doc.add_paragraph()
                self._add_word_text(p, label, size=12, bold=True)
                self._append_word_entries(p, examples)
            return
        rows = (len(entries) + 1) // 2
        if rows == 0:
            return
        table = doc.add_table(rows=rows, cols=2)
        table.style = "Table Grid"
        for idx, (label, examples) in enumerate(entries):
            r = idx % rows
            c = idx // rows
            cell = table.cell(r, c)
            para = cell.paragraphs[0]
            self._add_word_text(para, label, size=11, bold=True)
            self._append_word_entries(para, examples)
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    if not paragraph.text.strip():
                        paragraph.text = ""

    def _build_rich_text_cell_value(
        self, tone_map: dict[str, list[ParsedPhonologyRow]]
    ):
        if not _OPENPYXL_RICH_TEXT_AVAILABLE:
            return None
        rich_text = CellRichText()
        is_first_tone = True
        for tone_label in self._sort_tones(set(tone_map.keys())):
            if not is_first_tone:
                self._append_xlsx_rich_text(
                    rich_text, " ", font_name="Times New Roman", size=11
                )
            self._append_xlsx_rich_text(
                rich_text, f"[{tone_label}]", font_name="Times New Roman", size=11
            )
            for row in tone_map[tone_label]:
                self._append_xlsx_mixed_text(rich_text, row.character, size=11)
                if row.note:
                    self._append_xlsx_mixed_text(rich_text, row.note, size=11, subscript=True)
            is_first_tone = False
        return rich_text

    def _append_xlsx_mixed_text(self, rich_text, text: str, size: int, subscript: bool = False):
        if not text:
            return
        current_buffer = ""
        current_cjk = self._contains_cjk(text[0])
        for ch in text:
            flag = self._contains_cjk(ch)
            if flag != current_cjk and current_buffer:
                self._append_xlsx_rich_text(
                    rich_text,
                    current_buffer,
                    font_name="宋体" if current_cjk else "Times New Roman",
                    size=size,
                    subscript=subscript,
                )
                current_buffer = ""
                current_cjk = flag
            current_buffer += ch
        if current_buffer:
            self._append_xlsx_rich_text(
                rich_text,
                current_buffer,
                font_name="宋体" if current_cjk else "Times New Roman",
                size=size,
                subscript=subscript,
            )

    def _append_xlsx_rich_text(
        self,
        rich_text,
        text: str,
        font_name: str,
        size: int,
        subscript: bool = False,
    ):
        if not text:
            return
        font = InlineFont(
            rFont=font_name,
            sz=size,
            vertAlign="subscript" if subscript else None,
        )
        rich_text.append(TextBlock(font, text))

    def _setup_word_document_styles(self, doc):
        if not hasattr(doc, "styles"):
            return
        try:
            from docx.oxml.ns import qn
            from docx.shared import Pt
        except Exception:
            return
        style = doc.styles["Normal"]
        style.font.name = "Times New Roman"
        style.font.size = Pt(12)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

    def _center_word_paragraph(self, paragraph):
        if not hasattr(paragraph, "alignment"):
            return
        try:
            from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
        except Exception:
            return
        paragraph.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    def _style_word_paragraph_runs(self, paragraph, size: int = 12, bold: bool = False):
        runs = getattr(paragraph, "runs", [])
        for run in runs:
            self._apply_word_run_style(run, size=size, bold=bold)

    def _add_word_text(
        self,
        paragraph,
        text: str,
        size: int = 12,
        bold: bool = False,
        subscript: bool = False,
    ):
        if not text:
            return
        if not hasattr(paragraph, "runs"):
            run = paragraph.add_run(text)
            run.bold = bold
            if subscript:
                run.font.subscript = True
            return
        current_buffer = ""
        current_cjk = self._contains_cjk(text[0])
        for ch in text:
            flag = self._contains_cjk(ch)
            if flag != current_cjk and current_buffer:
                run = paragraph.add_run(current_buffer)
                self._apply_word_run_style(
                    run,
                    size=size,
                    bold=bold,
                    subscript=subscript,
                    cjk=current_cjk,
                )
                current_buffer = ""
                current_cjk = flag
            current_buffer += ch
        if current_buffer:
            run = paragraph.add_run(current_buffer)
            self._apply_word_run_style(
                run,
                size=size,
                bold=bold,
                subscript=subscript,
                cjk=current_cjk,
            )

    def _apply_word_run_style(
        self,
        run,
        size: int = 12,
        bold: bool = False,
        subscript: bool = False,
        cjk: bool = False,
    ):
        if not hasattr(run, "font"):
            return
        try:
            from docx.oxml.ns import qn
            from docx.shared import Pt
        except Exception:
            return
        run.bold = bold
        run.font.size = Pt(size)
        run.font.subscript = subscript
        run.font.name = "宋体" if cjk else "Times New Roman"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

    def _contains_cjk(self, text: str) -> bool:
        for ch in text:
            code = ord(ch)
            if 0x4E00 <= code <= 0x9FFF:
                return True
            if 0x3400 <= code <= 0x4DBF:
                return True
            if 0x20000 <= code <= 0x2A6DF:
                return True
        return False

    def _render_plain_entry(self, row: ParsedPhonologyRow) -> str:
        if not row.note:
            return row.character
        return f"{row.character}{self._to_subscript_text(row.note)}"

    def _to_subscript_text(self, value: str) -> str:
        mapping = str.maketrans(
            {
                "0": "₀",
                "1": "₁",
                "2": "₂",
                "3": "₃",
                "4": "₄",
                "5": "₅",
                "6": "₆",
                "7": "₇",
                "8": "₈",
                "9": "₉",
                "(": "₍",
                ")": "₎",
                "+": "₊",
                "-": "₋",
                "=": "₌",
                "a": "ₐ",
                "e": "ₑ",
                "o": "ₒ",
                "x": "ₓ",
                "ə": "ₔ",
            }
        )
        return value.translate(mapping)
