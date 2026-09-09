"""M01-C source_ids M01-PY-OPENPYXL, M01-PY-ET-XMLFILE, P06-SQLITE.

Construct and independently read back both artifacts before returning any output.
Public entry runs in an owned memory-limited Windows job. No existing database,
caller-selected output path, or unbounded worksheet temporary file is opened.
Durable publication/fencing is supplied by M01-F, not by this preparation layer.
"""
from dataclasses import asdict, dataclass
import io
import json
import math
import sqlite3
import struct
import sys
from zipfile import ZipFile, ZIP_DEFLATED
from .limits import Limits, LimitedBuffer, LimitError, FormatError, Cancelled
from .scratch import Scratch


@dataclass(frozen=True)
class ExportPair:
    xlsx: bytes
    sqlite: bytes
    columns: tuple[str,...]
    row_count: int


def checked_text(value):
    if type(value)!=str or len(value)>32767 or any(ord(c)<32 and c not in '\t\n\r' for c in value):
        raise FormatError('Invalid or excessive spreadsheet text')
    # Excel/XML forbids lone surrogate code points and FFFE/FFFF.
    if any(0xd800<=ord(c)<=0xdfff or ord(c) in (0xfffe,0xffff) for c in value):
        raise FormatError('Invalid spreadsheet Unicode')
    return value


def validate_table(table,limits):
    if type(table)!=dict or set(table)!={'columns','kinds','rows'}:raise FormatError('Invalid export table')
    columns,kinds,rows=(table[k] for k in ('columns','kinds','rows'))
    if not all(type(v)==list for v in (columns,kinds,rows)):raise FormatError('Invalid export vectors')
    if not columns or len(columns)>256 or len(columns)!=len(kinds):raise LimitError('export_column_limit')
    if (len(rows)+1)*len(columns)>limits.cells:raise LimitError('export_cell_limit')
    if not rows:raise FormatError('Empty analysis has no publishable parameter pair')
    text_size=0
    for c in columns:text_size+=len(checked_text(c).encode('utf-8'))
    if len({c.casefold() for c in columns})!=len(columns) or 'Time_s' not in columns:
        raise FormatError('Unique columns and Time_s are required')
    if any(k not in ('number','text') for k in kinds):raise FormatError('Invalid export column kind')
    ti=columns.index('Time_s');last=-math.inf
    if kinds[ti]!='number':raise FormatError('Time_s must be numeric')
    for row in rows:
        if type(row)!=list or len(row)!=len(columns):raise FormatError('Export row width mismatch')
        for k,value in zip(kinds,row):
            if k=='text':
                if value is not None:text_size+=len(checked_text(value).encode('utf-8'))
            elif value is not None and value not in ('+Infinity','-Infinity'):
                if type(value) not in (int,float) or not math.isfinite(value):raise FormatError('Invalid numeric export value')
            if text_size>limits.text_bytes:raise LimitError('export_text_limit')
        t=row[ti]
        if type(t) not in (int,float) or not math.isfinite(t) or t<last:raise FormatError('Invalid or unsorted Time_s')
        last=t
    return table


def table_from_frame(frame,limits):
    if (len(frame)+1)*len(frame.columns)>limits.cells:raise LimitError('export_cell_limit')
    columns=list(frame.columns)
    kinds=['number' if dtype.kind in 'biuf' else 'text' for dtype in frame.dtypes]
    rows=[]
    for raw in frame.itertuples(index=False,name=None):
        row=[]
        for kind,value in zip(kinds,raw):
            if kind=='number':
                value=float(value)
                value=None if math.isnan(value) else value if math.isfinite(value) else '+Infinity' if value>0 else '-Infinity'
            elif value is not None and not isinstance(value,str):
                # pandas represents an absent annotation as NaN; never stringify objects.
                if isinstance(value,float) and math.isnan(value):value=None
                else:raise FormatError('Non-text annotation')
            row.append(value)
        rows.append(row)
    return validate_table({'columns':columns,'kinds':kinds,'rows':rows},limits)


def _xlsx(table,limits):
    from openpyxl import Workbook
    from openpyxl.writer.excel import ExcelWriter
    from openpyxl.worksheet._writer import WorksheetWriter
    class MemoryWriter(ExcelWriter):
        def write_worksheet(self,ws):
            xml=LimitedBuffer(limits.xml_bytes)
            writer=WorksheetWriter(ws,out=xml)
            try:
                writer.write();ws._rels=writer._rels
                self._archive.writestr(ws.path[1:],xml.getvalue())
                self.manifest.append(ws)
            finally:
                writer.close();xml.close()
    wb=Workbook();ws=wb.active;ws.title='Sheet1'
    for r,values in enumerate([table['columns'],*table['rows']],1):
        for c,value in enumerate(values,1):
            if r>1 and table['kinds'][c-1]=='number' and isinstance(value,str):
                value='inf' if value=='+Infinity' else '-inf'
            cell=ws.cell(r,c,value)
            if isinstance(value,str):cell.data_type='s'  # =, +, -, @ labels are literal text
    output=LimitedBuffer(limits.output_bytes)
    with ZipFile(output,'w',ZIP_DEFLATED,allowZip64=False) as archive:
        MemoryWriter(wb,archive).write_data()
    return output.getvalue()


def _sqlite(table,limits,available):
    if available<12288:raise LimitError('sqlite_pair_budget_exceeded')
    conn=sqlite3.connect(':memory:')
    try:
        conn.execute('PRAGMA page_size=4096')
        conn.execute('PRAGMA temp_store=MEMORY')
        conn.execute('PRAGMA max_page_count='+str(available//4096))
        conn.setlimit(sqlite3.SQLITE_LIMIT_LENGTH,min(limits.text_bytes,available))
        conn.setlimit(sqlite3.SQLITE_LIMIT_COLUMN,256)
        columns=table['columns']
        quote=lambda name:'"'+name.replace('"','""')+'"'
        ddl=','.join(quote(c)+(' REAL' if k=='number' else ' TEXT') for c,k in zip(columns,table['kinds']))
        conn.execute('CREATE TABLE params ('+ddl+')')
        def rows():
            for row in table['rows']:
                yield tuple((float('inf') if v=='+Infinity' else float('-inf') if v=='-Infinity' else v) if k=='number' else v
                            for k,v in zip(table['kinds'],row))
        conn.executemany('INSERT INTO params VALUES ('+','.join('?' for _ in columns)+')',rows())
        conn.execute('CREATE INDEX idx_params_time ON params(Time_s)');conn.commit()
        blob=conn.serialize()
        if len(blob)>available:raise LimitError('sqlite_pair_budget_exceeded')
        return blob
    except sqlite3.DatabaseError as exc:
        raise FormatError('SQLite output could not be completed') from exc
    finally:conn.close()


def _equal(expected,actual,kind,excel=False):
    if kind=='text':return (expected or '')==(actual or '')
    if expected is None:return actual is None
    if expected in ('+Infinity','-Infinity'):
        target=(float('inf') if expected=='+Infinity' else float('-inf'))
        return actual==('inf' if target>0 else '-inf') if excel else actual==target
    return type(actual) in (int,float) and math.isclose(expected,actual,rel_tol=1e-12,abs_tol=1e-12)


def verify_pair(pair,table):
    """Read actual formats independently; headers, text types, every cell and index."""
    from openpyxl import load_workbook
    wb=load_workbook(io.BytesIO(pair.xlsx),read_only=True,data_only=False)
    try:
        rows=iter(wb.active.iter_rows())
        if [c.value for c in next(rows)]!=table['columns']:raise FormatError('XLSX header mismatch')
        count=0
        for expected in table['rows']:
            cells=next(rows,None)
            if cells is None or len(cells)!=len(expected):raise FormatError('XLSX shape mismatch')
            for kind,e,cell in zip(table['kinds'],expected,cells):
                if cell.data_type=='f' or not _equal(e,cell.value,kind,True):raise FormatError('XLSX cell mismatch')
            count+=1
        if next(rows,None) is not None:raise FormatError('XLSX extra rows')
    finally:wb.close()
    conn=sqlite3.connect(':memory:')
    try:
        conn.deserialize(pair.sqlite)
        if conn.execute('PRAGMA integrity_check').fetchone()!=('ok',):raise FormatError('SQLite integrity mismatch')
        if [r[1] for r in conn.execute('PRAGMA table_info(params)')]!=table['columns']:raise FormatError('SQLite header mismatch')
        if not any(r[1]=='idx_params_time' for r in conn.execute('PRAGMA index_list(params)')):raise FormatError('SQLite time index missing')
        rows=iter(conn.execute('SELECT * FROM params ORDER BY rowid'))
        for expected in table['rows']:
            actual=next(rows,None)
            if actual is None or any(not _equal(e,a,k) for k,e,a in zip(table['kinds'],expected,actual)):raise FormatError('SQLite cell mismatch')
        if next(rows,None) is not None:raise FormatError('SQLite extra rows')
    finally:conn.close()


def _build_pair(table,limits):
    """Internal worker implementation; call export_pair for process memory bounds."""
    validate_table(table,limits)
    xlsx=_xlsx(table,limits)
    sqlite=_sqlite(table,limits,limits.output_bytes-len(xlsx)-16)
    pair=ExportPair(xlsx,sqlite,tuple(table['columns']),len(table['rows']))
    verify_pair(pair,table)
    return pair


def export_pair(frame,scratch,limits=Limits(),stop=lambda:False,on_started=None):
    """Prepare a dataframe whose columns already have their intended export names."""
    if not isinstance(scratch,Scratch):raise TypeError('Host-owned Scratch capability required')
    if stop():raise Cancelled('cancelled')
    table=table_from_frame(frame,limits)
    request={'table':table,'limits':asdict(limits)}
    stream=LimitedBuffer(limits.input_bytes)
    for fragment in json.JSONEncoder(ensure_ascii=False,allow_nan=False,separators=(',',':')).iterencode(request):
        if stop():raise Cancelled('cancelled')
        stream.write(fragment.encode('utf-8'))
    path=scratch.create(stream.getvalue(),'.json')
    from ..native.windows import InputPipe
    from ..native.reaper import collect_pipe
    pipe=None
    try:
        pipe=InputPipe()
        payload,_=collect_pipe([sys.executable,'-B','-m','ptb_worker.io.export_worker',str(path),pipe.name],pipe,scratch.root,limits,stop,on_started)
        if len(payload)<16:raise FormatError('Incomplete export pair')
        nx,ns=struct.unpack_from('<QQ',payload)
        if 16+nx+ns!=len(payload):raise FormatError('Incomplete export pair')
        return ExportPair(payload[16:16+nx],payload[16+nx:],tuple(table['columns']),len(table['rows']))
    finally:
        if pipe:pipe.close()
        scratch.remove(path)


def export_analysis(result,scratch,limits=Limits(),stop=lambda:False,on_started=None):
    """Preserve the historical scientific key -> displayed export header mapping."""
    from phonetic_core.acoustic.catalog import PARAMETER_MAPPING
    return export_pair(result.to_dataframe().rename(columns=PARAMETER_MAPPING),scratch,limits,stop,on_started)
