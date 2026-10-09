"""Versioned M01 bundle, incremental disk writer and bounded interval reader.

Only application-defined ordinary tables are accepted. No user SQL, formulas,
extensions or schema functions are executed. Legacy imports keep their reader.
"""
import hashlib
import json
import math
import sqlite3
from pathlib import Path
import numpy as np
from phonetic_core.catalog import PARAMETER_MAPPING

VERSION='m01-bundle/2'
MAX_BYTES=4_000_000_000
MAX_CELLS=180_000_000

def label(key):
    if key=='gF0':return 'F0 - GCI'
    if key.endswith('_gF0'):
        return PARAMETER_MAPPING.get(key[:-4]+'_pF0',key).replace('(pF0)','(gF0)')
    return PARAMETER_MAPPING.get(key,key)

def quote(name):return '"'+name.replace('"','""')+'"'

def scalar(value):
    if isinstance(value,np.generic):value=value.item()
    if isinstance(value,float) and not math.isfinite(value):return None if math.isnan(value) else '+Infinity' if value>0 else '-Infinity'
    if value is None or type(value) in (float,int,str):return value
    raise ValueError('m01_invalid_output')

def digest(path):
    h=hashlib.sha256()
    with open(path,'rb') as f:
        for raw in iter(lambda:f.read(1048576),b''):h.update(raw)
    return h.hexdigest()

class BundleWriter:
    def __init__(self,path):
        self.path=Path(path)
        if self.path.exists():raise ValueError('Output already exists')
        self.conn=sqlite3.connect(self.path)
        self.conn.execute('PRAGMA journal_mode=DELETE')
        self.conn.execute('PRAGMA temp_store=FILE')
        self.conn.execute('PRAGMA cache_size=-8192')
        self.conn.execute('CREATE TABLE ptb_metadata (key TEXT PRIMARY KEY, value TEXT NOT NULL)')
        self.tables={};self.cells=0

    def append(self,name,arrays):
        if name not in ('params','egg_cycles'):raise ValueError('Invalid table')
        keys=list(arrays);n=len(arrays[keys[0]])
        columns=[label(k) if name=='params' else k for k in keys]
        if len(set(columns))!=len(columns) or any(len(arrays[k])!=n for k in keys):raise ValueError('m01_invalid_output')
        kinds=['text' if np.asarray(arrays[k]).dtype.kind in 'OUS' else 'number' for k in keys]
        if name not in self.tables:
            if not 1<=len(keys)<=256 or 'Time_s' not in keys:raise ValueError('m01_invalid_output')
            self.conn.execute('CREATE TABLE '+name+' ('+','.join(quote(c)+(' TEXT' if k=='text' else ' REAL') for c,k in zip(columns,kinds))+')')
            self.tables[name]=dict(columns=columns,kinds=kinds,rows=0,last=None)
        info=self.tables[name]
        if info['columns']!=columns:raise ValueError('m01_column_changed')
        if info['rows']+n>2_000_000:raise ValueError('m01_output_budget')
        times=np.asarray(arrays['Time_s'])
        if n and (not np.isfinite(times).all() or np.any(np.diff(times)<=0) or (info['last'] is not None and times[0]<=info['last'])):raise ValueError('m01_time_order')
        self.cells+=n*len(keys)
        if self.cells>MAX_CELLS:raise ValueError('m01_output_budget')
        self.conn.executemany('INSERT INTO '+name+' VALUES ('+','.join('?' for _ in keys)+')',
            (tuple(scalar(v) for v in row) for row in zip(*(arrays[k] for k in keys))))
        info['rows']+=n
        if n:info['last']=float(times[-1])
        self.conn.commit()
        if self.path.stat().st_size>MAX_BYTES:raise ValueError('m01_output_budget')

    def finish(self,metadata):
        for name in self.tables:self.conn.execute('CREATE INDEX '+name+'_time ON '+name+' (Time_s)')
        if 'F0_Time_s' in self.tables.get('egg_cycles',{}).get('columns',[]):
            self.conn.execute('CREATE INDEX egg_cycles_f0_time ON egg_cycles (F0_Time_s)')
        metadata={**metadata,'format_revision':VERSION,'tables':self.tables}
        self.conn.execute('INSERT INTO ptb_metadata VALUES (?,?)',('manifest',json.dumps(metadata,ensure_ascii=False,allow_nan=False)))
        self.conn.commit();self.conn.close()
        return metadata

def export_xlsx(sqlite_path,path,progress=lambda *a:None,check=lambda:None):
    from openpyxl import Workbook
    conn=sqlite3.connect('file:'+Path(sqlite_path).as_posix()+'?mode=ro',uri=True)
    wb=Workbook(write_only=True)
    try:
        meta=json.loads(conn.execute("SELECT value FROM ptb_metadata WHERE key='manifest'").fetchone()[0])
        total=sum(i['rows'] for i in meta['tables'].values());done=0
        for table,info in meta['tables'].items():
            sheet=None
            for i,row in enumerate(conn.execute('SELECT * FROM '+table+' ORDER BY Time_s')):
                if i%1_000_000==0:
                    sheet=wb.create_sheet(table if i==0 else table+'_'+str(i//1_000_000+1));sheet.append(info['columns'])
                # Explicit text cells prevent annotations beginning '=' becoming formulas.
                from openpyxl.cell import WriteOnlyCell
                cells=[]
                for value in row:
                    cell=WriteOnlyCell(sheet,value=value)
                    if isinstance(value,str):cell.data_type='s'
                    cells.append(cell)
                sheet.append(cells);done+=1
                if done%2000==0:check();progress('xlsx',done/max(total,1))
            if sheet is None:wb.create_sheet(table).append(info['columns'])
        ws=wb.create_sheet('ptb_metadata');ws.append(['format_revision',VERSION])
        # Chunk metadata to remain inside Excel's per-cell text bound.
        text=json.dumps(meta,ensure_ascii=False,allow_nan=False)
        for offset in range(0,len(text),16000):ws.append(['manifest',text[offset:offset+16000]])
        check();wb.save(path)
    finally:conn.close();wb.close()

def open_checked(path):
    path=Path(path)
    if not 0<path.stat().st_size<=MAX_BYTES:raise ValueError('parameter_input_budget')
    conn=sqlite3.connect(path.as_uri()+'?mode=ro',uri=True)
    try:
        conn.execute('PRAGMA trusted_schema=OFF');conn.execute('PRAGMA query_only=ON');conn.execute('PRAGMA cache_size=-8192')
        conn.setlimit(sqlite3.SQLITE_LIMIT_LENGTH,2_000_000)
        conn.setlimit(sqlite3.SQLITE_LIMIT_COLUMN,256)
        conn.setlimit(sqlite3.SQLITE_LIMIT_SQL_LENGTH,65536)
        remaining=[MAX_CELLS*50]
        def budget():
            remaining[0]-=10000
            return remaining[0]<0
        conn.set_progress_handler(budget,10000)
        schema=conn.execute("SELECT type,name,sql FROM sqlite_schema WHERE name NOT LIKE 'sqlite_%'").fetchall()
        if len(schema)>8 or any(t not in ('table','index') or (t=='table' and (n not in ('params','egg_cycles','ptb_metadata') or not s or 'VIRTUAL' in s.upper() or 'GENERATED' in s.upper())) for t,n,s in schema):raise ValueError('invalid_parameter_table')
        if not {'params','ptb_metadata'} <= {n for t,n,_ in schema if t=='table'}:raise ValueError('not_parameter_bundle')
        row=conn.execute("SELECT value FROM ptb_metadata WHERE key='manifest' LIMIT 1").fetchone()
        meta=json.loads(row[0])
        if meta['format_revision']!=VERSION:raise ValueError('not_parameter_bundle')
        for name,info in meta['tables'].items():
            if name not in ('params','egg_cycles'):raise ValueError('invalid_parameter_table')
            fields=conn.execute('PRAGMA table_xinfo('+name+')').fetchall()
            if any(r[6]!=0 or r[1].lower() in ('rowid','_rowid_','oid') for r in fields):raise ValueError('invalid_parameter_table')
            columns=[r[1] for r in fields]
            if columns!=info['columns'] or len(columns)!=len(info['kinds']) or 'Time_s' not in columns:raise ValueError('invalid_parameter_table')
            if type(info['rows'])!=int or not 0<=info['rows']<=2_000_000 or any(k not in ('number','text') for k in info['kinds']):raise ValueError('invalid_parameter_table')
        if not 0<float(meta['duration_s'])<=1800:raise ValueError('invalid_parameter_table')
        if sum(i['rows']*len(i['columns']) for i in meta['tables'].values())>MAX_CELLS:raise ValueError('parameter_input_budget')
        return conn,meta
    except BaseException:conn.close();raise

def window(path,sha,view=None):
    conn,meta=open_checked(Path(path).absolute())
    try:
        view=view or {};start=float(view.get('start',0));end=float(view.get('end',meta['duration_s']))
        width=max(32,min(4000,int(view.get('width',1200))))
        if not math.isfinite(start+end) or not 0<=start<end<=max(1800.,meta['duration_s']):raise ValueError('invalid_parameter_range')
        catalog=[c for c in meta['tables']['params']['columns'] if c!='Time_s']
        native=meta.get('egg') and meta['egg']['storage']=='cycles'
        if native:catalog += ['CQ','SQ','F0 - GCI']
        names=view.get('parameters',[])
        if not isinstance(names,list) or len(names)>128 or any(n not in catalog for n in names):raise ValueError('invalid_parameter_columns')
        tracks={};stats={}
        for name in names:
            check_name='gF0' if name=='F0 - GCI' else name
            table='egg_cycles' if native and name in ('CQ','SQ','F0 - GCI') else 'params'
            column=check_name if table=='egg_cycles' else name
            time_column='F0_Time_s' if table=='egg_cycles' and name=='F0 - GCI' else 'Time_s'
            bucket=[];last_pixel=None;points=[];count=0;total=0.;lo=None;hi=None
            def flush():
                if not bucket:return
                finite=[(i,p) for i,p in enumerate(bucket) if isinstance(p[1],(int,float))]
                indexes={0,len(bucket)-1}
                gap=next((i for i,p in enumerate(bucket) if p[1] is None),None)
                if gap is not None:indexes.add(gap)
                if finite:indexes.update((min(finite,key=lambda a:a[1][1])[0],max(finite,key=lambda a:a[1][1])[0]))
                points.extend(bucket[i] for i in sorted(indexes));bucket.clear()
            for t,v in conn.execute('SELECT '+quote(time_column)+','+quote(column)+' FROM '+table+' WHERE '+quote(time_column)+'>=? AND '+quote(time_column)+'<=? ORDER BY '+quote(time_column),(start,end)):
                if t<start or t>end:continue
                if isinstance(v,(int,float)) and math.isfinite(v):
                    count+=1;total+=abs(v);lo=v if lo is None else min(lo,v);hi=v if hi is None else max(hi,v)
                pixel=int((t-start)/(end-start)*width+1e-9)
                if pixel!=last_pixel:flush();last_pixel=pixel
                bucket.append([t,v])
            flush();tracks[name]=points;stats[name]=dict(count=count,mean=total/count if count else 0,min=lo,max=hi)
        return dict(schema_version='m02/1',sha256=sha,columns=['Time_s',*catalog],
                    kinds=['number',*[('text' if n.startswith('text_') else 'number') for n in catalog]],rows=[],
                    tracks=tracks,stats=stats,streamed=True,duration_s=meta['duration_s'],metadata=meta)
    finally:conn.close()
