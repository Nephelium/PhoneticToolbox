"""Read old M01 XLSX/SQLite bytes inside the owned segmentation child.

source_ids: M01-PY-OPENPYXL, P06-SQLITE. Never opens the user's database path,
evaluates formulas, follows links, repairs rows, or claims original WAV identity.
"""
import io
import math
import re
import sqlite3
from zipfile import ZipFile,BadZipFile
from xml.etree import ElementTree as ET
from .limits import Limits,FormatError,LimitError
from .parameter_exports import validate_table

SUFFIXES=('.xlsx','.ptb.sqlite','.ptb.sqlite3')


def _xlsx(raw,limits):
    from openpyxl import load_workbook
    with ZipFile(io.BytesIO(raw)) as archive:
        entries=archive.infolist()
        if len(entries)>128 or sum(v.file_size for v in entries)>limits.xml_bytes:raise LimitError('legacy_xml_budget')
        if len({v.filename for v in entries})!=len(entries):raise FormatError('Duplicate ZIP member')
        for entry in entries:
            if entry.flag_bits&1:raise FormatError('Encrypted workbook')
            if entry.filename.endswith(('.xml','.rels')):
                data=archive.read(entry)
                # XLSX written by v2 uses UTF-8 XML; reject alternate encodings
                # before entity inspection so UTF-16 cannot bypass this check.
                data.decode('utf-8-sig')
                if b'\x00' in data:raise FormatError('Unsupported XML encoding')
                if b'<!DOCTYPE' in data or b'<!ENTITY' in data:raise FormatError('XML entities are unsupported')
                if b'externalLink' in data or re.search(rb'TargetMode\s*=\s*[\x22\x27]External',data):raise FormatError('External workbook links are unsupported')
                depth=0;nodes=0
                for event,element in ET.iterparse(io.BytesIO(data),events=('start','end')):
                    if event=='start':
                        depth+=1;nodes+=1
                        if depth>32 or nodes>limits.cells*10+10000:raise LimitError('legacy_xml_structure_budget')
                        if element.tag.rsplit('}',1)[-1]=='c':
                            coordinate=element.get('r','');match=re.fullmatch(r'([A-Z]{1,3})([1-9][0-9]{0,6})',coordinate)
                            if not match:raise FormatError('Invalid worksheet cell coordinate')
                            column=0
                            for char in match[1]:column=column*26+ord(char)-64
                            if column>256 or int(match[2])>limits.cells:raise LimitError('legacy_worksheet_extent_budget')
                        if element.tag.rsplit('}',1)[-1]=='row':
                            index=element.get('r','')
                            if not index.isascii() or not index.isdigit() or len(index)>7 or not 1<=int(index)<=limits.cells:raise LimitError('legacy_worksheet_extent_budget')
                    else:depth-=1;element.clear()
    wb=load_workbook(io.BytesIO(raw),read_only=True,data_only=False,keep_links=False)
    try:
        if len(wb.worksheets)!=1:raise FormatError('Expected one parameter worksheet')
        sheet=wb.worksheets[0];sheet.reset_dimensions();rows=[];used=0
        for cells in sheet.iter_rows():
            used+=len(cells)
            if used>limits.cells:raise LimitError('legacy_cell_budget')
            if any(c.data_type in ('f','e') for c in cells):raise FormatError('Formula or error cell in legacy table')
            rows.append([c.value for c in cells])
        if not rows:raise FormatError('Empty parameter worksheet')
        width=len(rows[0])
        if any(len(r)>width for r in rows):raise FormatError('Unlabelled parameter column')
        return rows[0],[r+[None]*(width-len(r)) for r in rows[1:]]
    finally:wb.close()


def _sqlite(raw,limits):
    if not raw.startswith(b'SQLite format 3\x00'):raise FormatError('Invalid SQLite header')
    conn=sqlite3.connect(':memory:')
    try:
        conn.setlimit(sqlite3.SQLITE_LIMIT_LENGTH,min(limits.input_bytes,limits.text_bytes))
        conn.setlimit(sqlite3.SQLITE_LIMIT_COLUMN,256)
        conn.setlimit(sqlite3.SQLITE_LIMIT_SQL_LENGTH,65536)
        conn.execute('PRAGMA temp_store=MEMORY');conn.execute('PRAGMA trusted_schema=OFF')
        conn.deserialize(raw);conn.execute('PRAGMA query_only=ON')
        remaining=[max(1000,limits.cells*20)]
        def progress():
            remaining[0]-=1000
            return remaining[0]<0
        conn.set_progress_handler(progress,1000)
        schema=conn.execute("SELECT type,name,sql FROM sqlite_schema WHERE name NOT LIKE 'sqlite_%' LIMIT 130").fetchall()
        if len(schema)>128 or any(t not in ('table','index') or (t=='table' and (n!='params' or not s or 'VIRTUAL' in s.upper() or 'GENERATED' in s.upper())) for t,n,s in schema):raise FormatError('Expected ordinary params table only')
        if sum(t=='table' and n=='params' for t,n,_ in schema)!=1:raise FormatError('Missing params table')
        info=conn.execute('PRAGMA table_xinfo(params)').fetchall()
        if any(r[6]!=0 or r[1].lower() in ('rowid','_rowid_','oid') for r in info):raise FormatError('Generated or shadowing columns are unsupported')
        columns=[r[1] for r in info]
        if not columns:raise FormatError('Empty params schema')
        # No schema functions, attach, writes, extension calls or other tables.
        def authorize(action,arg1,arg2,db,trigger):
            return sqlite3.SQLITE_OK if action==sqlite3.SQLITE_SELECT or (action==sqlite3.SQLITE_READ and arg1=='params' and db=='main' and trigger is None) else sqlite3.SQLITE_DENY
        conn.set_authorizer(authorize)
        rows=[]
        for row in conn.execute('SELECT * FROM params ORDER BY rowid'):
            if (len(rows)+2)*len(columns)>limits.cells:raise LimitError('legacy_cell_budget')
            rows.append(list(row))
        return columns,rows
    except sqlite3.DatabaseError:raise FormatError('Invalid or excessive legacy SQLite') from None
    finally:conn.close()


def read_legacy_parameters(raw,name,limits=Limits()):
    if type(raw)!=bytes or not 0<len(raw)<=limits.input_bytes:raise LimitError('legacy_input_budget')
    if type(name)!=str or not name.lower().endswith(SUFFIXES):raise FormatError('Unsupported parameter format')
    try:columns,rows=_xlsx(raw,limits) if name.lower().endswith('.xlsx') else _sqlite(raw,limits)
    except (BadZipFile,ET.ParseError,KeyError,ValueError,TypeError,OverflowError) as exc:
        if isinstance(exc,(FormatError,LimitError)):raise
        raise FormatError('Invalid legacy parameter table') from None
    if any(type(c)!=str or not c for c in columns):raise FormatError('Parameter headers must be nonempty text')
    kinds=[]
    for index,column in enumerate(columns):
        values=[r[index] for r in rows]
        annotation=column.lower()=='textgrid' or column.lower().startswith(('textgrid_','text_'))
        numeric=not annotation and all(v is None or type(v) in (int,float) or (type(v)==str and v.lower() in ('nan','inf','+inf','-inf','+infinity','-infinity')) for v in values)
        kinds.append('number' if numeric else 'text')
        for row in rows:
            value=row[index]
            if numeric and value is not None:
                value=float(value)
                row[index]=None if math.isnan(value) else value if math.isfinite(value) else '+Infinity' if value>0 else '-Infinity'
            elif annotation and value is None:row[index]=''
    return validate_table({'columns':columns,'kinds':kinds,'rows':rows},limits)
