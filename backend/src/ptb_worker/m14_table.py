"""M14/2 bounded table decoding; preserves empty columns and source row numbers."""
import csv
import io
from pathlib import PurePath
from zipfile import ZipFile
import re

from phonetic_core.transcription.phonology import PhonologyRules
from .m14_import import MAX_INPUT, MAX_ROWS, MAX_CELL

EXTENSIONS = ('.xlsx', '.xls', '.csv', '.tsv', '.txt', '.docx')


def _zip_check(raw, suffix):
    with ZipFile(io.BytesIO(raw)) as archive:
        entries = archive.infolist()
        if len(entries) > 512 or sum(e.file_size for e in entries) > 32_000_000:
            raise ValueError('m14_expanded_budget')
        if suffix == '.xlsx':
            for entry in entries:
                if entry.filename.startswith('xl/worksheets/') and entry.filename.endswith('.xml'):
                    if re.search(rb'<(?:\w+:)?f(?:\s|>)', archive.read(entry)):
                        raise ValueError('m14_formula_input')


def read_table(raw, name, options):
    if not isinstance(raw, bytes) or not raw:
        raise ValueError('m14_empty_input')
    if len(raw) > MAX_INPUT:
        raise ValueError('m14_input_budget')
    suffix = PurePath(name).suffix.lower()
    if suffix not in EXTENSIONS:
        raise ValueError('m14_unsupported_format')
    index = options.get('table_index', 0)
    if type(index) is not int or index < 0:
        raise ValueError('m14_table_index')
    tables = []; records = []
    try:
        if suffix in ('.xlsx', '.docx'):
            _zip_check(raw, suffix)
        if suffix == '.xlsx':
            from openpyxl import load_workbook
            book = load_workbook(io.BytesIO(raw), read_only=True, data_only=False)
            try:
                tables = book.sheetnames
                if index >= len(tables): raise ValueError('m14_table_index')
                sheet = book[tables[index]]
                if sheet.max_column and sheet.max_column > 64: raise ValueError('m14_column_budget')
                for number, row in enumerate(sheet.iter_rows(values_only=True), 1):
                    if number > MAX_ROWS + 1: raise ValueError('m14_row_budget')
                    records.append((number, ['' if v is None else str(v) for v in row]))
            finally: book.close()
        elif suffix == '.xls':
            import xlrd
            book = xlrd.open_workbook(file_contents=raw, on_demand=True)
            try:
                tables = book.sheet_names()
                if index >= len(tables): raise ValueError('m14_table_index')
                sheet = book.sheet_by_index(index)
                if sheet.nrows > MAX_ROWS + 1: raise ValueError('m14_row_budget')
                if sheet.ncols > 64: raise ValueError('m14_column_budget')
                for r in range(sheet.nrows):
                    records.append((r + 1, [str(int(v)) if isinstance(v, float) and v.is_integer() else str(v) for v in sheet.row_values(r)]))
            finally: book.release_resources()
        elif suffix == '.docx':
            from docx import Document
            book = Document(io.BytesIO(raw))
            tables = [f'表格 {i + 1}' for i in range(len(book.tables))]
            if index >= len(tables): raise ValueError('m14_table_index')
            for number, row in enumerate(book.tables[index].rows, 1):
                if number > MAX_ROWS + 1: raise ValueError('m14_row_budget')
                records.append((number, [c.text for c in row.cells]))
        else:
            if index: raise ValueError('m14_table_index')
            tables = ['文本表格']
            encoding = options.get('encoding', 'auto')
            if encoding == 'auto': encoding = 'utf-16' if raw.startswith((b'\xff\xfe', b'\xfe\xff')) else 'utf-8-sig'
            if encoding not in ('utf-8-sig', 'utf-8', 'gb18030', 'utf-16'): raise ValueError('m14_encoding')
            text = raw.decode(encoding)
            delimiter = options.get('delimiter', 'auto')
            delimiters = {'tab': '\t', 'comma': ',', 'semicolon': ';', 'chinese_comma': '，', 'space': ' '}
            if delimiter == 'auto':
                if suffix == '.tsv': char = '\t'
                else:
                    try: char = csv.Sniffer().sniff(text[:8192], delimiters=',;\t，').delimiter
                    except csv.Error: char = ',' if suffix == '.csv' else '\t'
            elif delimiter in delimiters: char = delimiters[delimiter]
            else: raise ValueError('m14_delimiter')
            reader = csv.reader(io.StringIO(text, newline=''), delimiter=char, strict=True)
            previous = 0
            for cols in reader:
                number = previous + 1; previous = reader.line_num
                if len(records) >= MAX_ROWS + 1: raise ValueError('m14_row_budget')
                if number < options.get('start_row', 2) <= previous:
                    raise ValueError('m14_start_inside_record')
                records.append((number, cols))
        width = max((len(cols) for _, cols in records), default=0)
        if width > 64: raise ValueError('m14_column_budget')
        for _, cols in records:
            if any(len(v) > MAX_CELL or any(ord(c) < 32 and c not in '\t\r\n' for c in v) for v in cols):
                raise ValueError('m14_invalid_cell')
        return records, tables, width
    except Exception as exc:
        if isinstance(exc, ValueError) and str(exc).startswith('m14_'): raise
        raise ValueError('m14_decode_failed') from exc


def inspect(raw, name, options):
    records, tables, width = read_table(raw, name, dict(options, start_row=1))
    start=options.get('start_row',2)
    sample=records[:6]+[r for r in records[6:] if r[0]>=start][:18]
    return dict(schema_version='m14-table/1', tables=tables, table_index=options.get('table_index', 0),
                column_count=width, total_rows=len(records), sample=[dict(row=n, cells=c) for n, c in sample])


def load_v2(raw, name, options):
    records, _, width = read_table(raw, name, options)
    char_col = options.get('character_column', 1); ipa_col = options.get('ipa_column', 2)
    note_col = options.get('note_column', 3); start = options.get('start_row', 2)
    if type(start) is not int or not 1 <= start <= MAX_ROWS + 1: raise ValueError('m14_start_row')
    chosen = [char_col, ipa_col] + ([] if note_col is None else [note_col])
    if any(type(v) is not int or not 1 <= v <= 64 for v in chosen) or len(set(chosen)) != len(chosen):
        raise ValueError('m14_column_selection')
    if max(char_col, ipa_col) > width: raise ValueError('m14_missing_columns')
    if note_col is not None and note_col > width and note_col != 3: raise ValueError('m14_missing_columns')
    rows = []; skipped = []; numbers = []; warnings = []
    rules = PhonologyRules(computation_revision='m14/2')
    for number, cols in records:
        if number < start:
            skipped.append(dict(row=number, reason='数据开始行之前')); continue
        def cell(c): return cols[c - 1] if c is not None and c <= len(cols) else ''
        selected = [cell(char_col), cell(ipa_col), cell(note_col)]
        row = rules._parse_columns(selected)
        if row is None:
            skipped.append(dict(row=number, reason='表头、空行或缺少字头 / IPA')); continue
        base = row.ipa.strip().strip('[]/').strip()
        if re.fullmatch(r'[0-9０-９⁰¹²³⁴⁵⁶⁷⁸⁹ˈˌ]+',base):
            skipped.append(dict(row=number,reason='只有调值或重音符号，缺少音段')); continue
        if re.search(r'\s', base) or '.' in base or re.search(r'[0-9０-９⁰¹²³⁴⁵⁶⁷⁸⁹]+[^0-9０-９⁰¹²³⁴⁵⁶⁷⁸⁹]', base):
            skipped.append(dict(row=number, reason='疑似多个音节，请将每个音节放在独立记录中')); continue
        if '?' in base or '�' in base:
            warnings.append(dict(row=number, reason='含未知或替换符号，请核对原始 IPA'))
        rows.append(row); numbers.append(number)
        if len(rows) > MAX_ROWS: raise ValueError('m14_row_budget')
    if not rows: raise ValueError('m14_no_valid_rows')
    duplicates = len(rows) - len(set((r.character, r.ipa, r.note) for r in rows))
    return rows, dict(total_rows=len(records), accepted_rows=len(rows), skipped=skipped, duplicate_rows=duplicates, source_rows=numbers, warnings=warnings)
