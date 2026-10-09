"""Owned file-backed parameter reader. XLSX is indexed once, then queried by view."""
import json
import sys
from pathlib import Path
from .acoustic_stream_child import atomic


def convert_xlsx(source,target):
    from zipfile import ZipFile
    from openpyxl import load_workbook
    import numpy as np
    from .io.parameter_bundle import BundleWriter,VERSION,MAX_CELLS
    with ZipFile(source) as archive:
        entries=archive.infolist()
        if len(entries)>128 or len({e.filename for e in entries})!=len(entries) or sum(e.file_size for e in entries)>16_000_000_000:
            raise ValueError('parameter_input_budget')
        for entry in entries:
            if entry.flag_bits&1:raise ValueError('invalid_parameter_table')
            if entry.filename.endswith(('.xml','.rels')):
                previous=b''
                with archive.open(entry) as stream:
                    for block in iter(lambda:stream.read(1048576),b''):
                        raw=previous+block
                        if any(token in raw for token in (b'\x00',b'<!DOCTYPE',b'<!ENTITY',b'externalLink',b'TargetMode="External"',b"TargetMode='External'")):
                            raise ValueError('invalid_parameter_table')
                        previous=raw[-128:]
        # openpyxl loads shared strings eagerly; our writer uses inline strings.
        strings=next((e for e in entries if e.filename=='xl/sharedStrings.xml'),None)
        if strings and strings.file_size>16_000_000:raise ValueError('parameter_input_budget')
    # Managed assets intentionally have a .bin extension. Validate ZIP contents,
    # then pass the granted stream so openpyxl does not infer format from .bin.
    handle=Path(source).open('rb')
    try:wb=load_workbook(handle,read_only=True,data_only=False,keep_links=False)
    except BaseException:handle.close();raise
    writer=None
    try:
        if 'ptb_metadata' not in wb.sheetnames:raise ValueError('not_parameter_bundle')
        parts=[];revision=None
        for i,row in enumerate(wb['ptb_metadata'].iter_rows(max_col=2,values_only=True)):
            if i>128:raise ValueError('invalid_parameter_table')
            if row[0]=='format_revision':revision=row[1]
            elif row[0]=='manifest' and isinstance(row[1],str):parts.append(row[1])
            else:raise ValueError('invalid_parameter_table')
        if revision!=VERSION:raise ValueError('not_parameter_bundle')
        metadata=json.loads(''.join(parts));infos=metadata['tables']
        if not {'params'}<=infos.keys() or any(n not in ('params','egg_cycles') for n in infos):raise ValueError('invalid_parameter_table')
        if sum(v['rows']*len(v['columns']) for v in infos.values())>MAX_CELLS:raise ValueError('parameter_input_budget')
        expected={'ptb_metadata'}
        for name,info in infos.items():
            expected.update(name if n==0 else name+'_'+str(n+1) for n in range(max(1,(info['rows']+999999)//1000000)))
        if set(wb.sheetnames)!=expected:raise ValueError('invalid_parameter_table')
        writer=BundleWriter(target)
        for name,info in infos.items():
            total=0
            for n in range(max(1,(info['rows']+999999)//1000000)):
                sheet=wb[name if n==0 else name+'_'+str(n+1)];sheet.reset_dimensions()
                rows=sheet.iter_rows(max_col=len(info['columns']));header=next(rows)
                if [c.value for c in header]!=info['columns']:raise ValueError('invalid_parameter_table')
                batch=[]
                def flush():
                    if batch:
                        arrays={c:np.array([r[i] for r in batch],dtype=object if info['kinds'][i]=='text' else float) for i,c in enumerate(info['columns'])}
                        writer.append(name,arrays);batch.clear()
                for cells in rows:
                    total+=1
                    if total>info['rows'] or len(cells)!=len(info['columns']) or any(c.data_type in ('f','e') for c in cells):raise ValueError('invalid_parameter_table')
                    batch.append([c.value for c in cells])
                    if len(batch)==2000:flush()
                flush()
            if total!=info['rows']:raise ValueError('invalid_parameter_table')
            if not total:writer.append(name,{c:np.array([],dtype=object if info['kinds'][i]=='text' else float) for i,c in enumerate(info['columns'])})
        writer.finish(metadata)
    finally:
        wb.close();handle.close()
        if writer:writer.conn.close()


def run(request):
    from .io.parameter_bundle import window
    source=Path(request['source']);cache=Path(request['cache']);name=request['name']
    try:
        if name.lower().endswith('.xlsx'):
            if not cache.exists():
                from uuid import uuid4
                temporary=cache.with_suffix('.'+uuid4().hex+'.building')
                # Each attempt has a fresh cache target; incomplete files are never read.
                convert_xlsx(source,temporary);temporary.replace(cache)
            source=cache
        return window(source,request['sha256'],request.get('view'))
    except ValueError as error:
        if str(error)!='not_parameter_bundle':raise
        if source.stat().st_size>16_000_000:raise ValueError('parameter_input_budget')
        from .io.legacy_parameters import read_legacy_parameters
        return dict(schema_version='m02/1',sha256=request['sha256'],**read_legacy_parameters(source.read_bytes(),name))


if __name__=='__main__':
    root=Path(sys.argv[1])
    try:value={'result':run(json.loads((root/'request.json').read_text('utf-8')))}
    except Exception as error:value={'error':str(error) if str(error).startswith(('parameter_','invalid_parameter','not_parameter')) else 'invalid_parameter_table'}
    atomic(root/'response.json',value)
