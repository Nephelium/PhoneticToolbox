import copy
import io
import json
from pathlib import Path
from zipfile import ZipFile
import pytest
from openpyxl import load_workbook,Workbook
from ptb_worker.m14_import import load
from ptb_worker.m14_jobs import execute
from phonetic_core.transcription.phonology import PhonologyRules,PhonologyInputRow
from phonetic_core.transcription.phonology.config import default_config,configure
from phonetic_core.transcription.phonology.export import export
from phonetic_core.transcription.phonology.presentation import WorkbenchRenderer

FONT=dict(schema_version='font/1',zh='Microsoft YaHei',latin='Segoe UI',ipa='Doulos SIL',size_px=14)
ROOT=Path(__file__).resolve().parents[2]


def test_fix01_three_outputs_preserve_configured_tone_order_and_duplicates():
    a=PhonologyRules().analyze([PhonologyInputRow('妈','ma55'),PhonologyInputRow('麻','ma35'),PhonologyInputRow('妈','ma55')])
    result=export(a,{'55':'乙','35':'甲'},['55','35'],['m'],['a'],renderer=WorkbenchRenderer(FONT))
    w=load_workbook(io.BytesIO(result['同音字表_二维表.xlsx']),rich_text=True)
    assert str(w.active['B2'].value)=='[乙]妈妈 [甲]麻'
    assert w.active.freeze_panes=='B2'
    assert w.active['B1'].font.name=='Doulos SIL'
    with ZipFile(io.BytesIO(result['同音字表_二维表.xlsx'])) as z:
        assert b'<t xml:space="preserve"> </t>' in z.read('xl/worksheets/sheet1.xml')
    for name,raw in result.items():
        if name.endswith('.docx'):
            from docx import Document
            d=Document(io.BytesIO(raw));text=''.join(p.text for p in d.paragraphs)
            assert text.index('[乙]')<text.index('[甲]')
            assert '妈妈' not in text and '妈 妈' in text
            with ZipFile(io.BytesIO(raw)) as z:assert 'Doulos SIL' in z.read('word/document.xml').decode()


def test_fix02_empty_final_merge_and_source_immutability():
    a=PhonologyRules().analyze([PhonologyInputRow('嗯','m35'),PhonologyInputRow('妈','ma55')],False)
    before=copy.deepcopy(a);s=default_config(a);s['final_map']={'':'a'};s['final_order']=['a']
    b,_=configure(a,s);assert a==before and len(b.rows)==2 and b.rows[0].ipa=='m35' and b.rows[0].final=='a'
    s['initial_map']={'m':'m'}
    with pytest.raises(ValueError,match='cyclic'):configure(a,s)


def test_invalid_order_and_unknown_mapping():
    a=PhonologyRules().analyze([PhonologyInputRow('妈','ma55')]);s=default_config(a);s['initial_order']=[]
    with pytest.raises(ValueError,match='order_mismatch'):configure(a,s)
    s=default_config(a);s['initial_map']={'p':'m'}
    with pytest.raises(ValueError,match='unknown_symbol'):configure(a,s)


@pytest.mark.parametrize('raw,name,code',[(b'x','a.wav','unsupported_format'),(b'\xff','a.txt','decode_failed'),(b'a'*2_000_001,'a.csv','input_budget'),(b'word\nonly','a.csv','missing_columns')],ids=['extension','encoding','size','columns'])
def test_decode_errors(raw,name,code):
    with pytest.raises(ValueError,match=code):load(raw,name)


def test_formula_and_zip_expansion_rejected():
    b=io.BytesIO();w=Workbook();w.active.append(['妈','=1+1']);w.save(b)
    with pytest.raises(ValueError,match='formula_input'):load(b.getvalue(),'a.xlsx',False)


def test_real_handler_all_formats_and_exact_three_outputs():
    for ext in ('xlsx','xls','csv','txt','tsv'):
        name='public.'+ext;raw=(ROOT/'tests/fixtures/m14'/name).read_bytes()
        config=dict(action='preview',skip_first_row=True,consonant_only_as_zero_initial=True)
        p=json.loads(execute(raw,name,config)['m14-preview.json'])
        out=execute(raw,name,dict(config,action='export',settings=p['config'],font=FONT))
        assert len(out)==3 and sum(n.endswith('.docx') for n in out)==2
        # V2 text splitting collapses empty fields: the third column becomes IPA.
        # Excel/CSV also interpret literal NA as missing; text preserves it.
        assert p['diagnostics']['accepted_rows']==(18 if ext in ('txt','tsv') else 16)


def test_cancelled_export_publishes_nothing():
    a=PhonologyRules().analyze([PhonologyInputRow('妈','ma55')]);calls=[]
    def cancelled():calls.append(1);return len(calls)>1
    with pytest.raises(InterruptedError):export(a,{},None,None,None,cancelled=cancelled)

def test_worker_dispatch_import_stays_light_before_heartbeat():
    import subprocess,sys
    code="import sys;import ptb_worker.m14_executor;assert not any(k in sys.modules for k in ('openpyxl','pandas','docx'))"
    subprocess.run([sys.executable,'-c',code],check=True,timeout=8)


def test_no_header_skip_actually_changes_count():
    raw='妈,ma55\n麻,ma35'.encode()
    rows,_=load(raw,'noheader.csv',False);skipped,_=load(raw,'noheader.csv',True)
    assert [r.character for r in rows]==['妈','麻']
    assert [r.character for r in skipped]==['麻']


def test_cell_overflow_rejected_without_truncating_records():
    a=PhonologyRules().analyze([PhonologyInputRow('妈','ma55','n'*512)]*70)
    with pytest.raises(ValueError,match='m14_cell_output_budget'):export(a,{},None,None,None)
    assert len(a.rows)==70 and all(len(r.note)==512 for r in a.rows)


def test_merged_output_counts_order_and_word_excel_consistency():
    from dataclasses import asdict
    baseline=json.loads((ROOT/'tests/fixtures/m14/v2-baseline.json').read_text(encoding='utf8'))['cases']['public.xlsx:True']['policies']['True']
    rows,_=load((ROOT/'tests/fixtures/m14/public.xlsx').read_bytes(),'public.xlsx');r=PhonologyRules();a=r.analyze(rows)
    b=r.apply_symbol_aliases(a,{'pʰ':'p','p':'m'},{'ã':'a'})
    assert asdict(b)==baseline['aliases'] and len(b.rows)==len(a.rows)==16
    assert [(x.character,x.ipa,x.note) for x in b.rows]==[(x.character,x.ipa,x.note) for x in a.rows]
    outputs=export(b,{x:x for x in b.unique_tones},baseline['tones'],b.unique_initials,b.unique_finals,renderer=WorkbenchRenderer(FONT))
    w=load_workbook(io.BytesIO(outputs['同音字表_二维表.xlsx']),rich_text=True).active
    column=[c.value for c in w[1]].index('m')+1
    row=next(i for i in range(2,w.max_row+1) if w.cell(i,1).value=='a')
    assert str(w.cell(row,column).value)=='[214]马1 [55]妈巴妈 [51]怕 [35]麻文鼻鼻化'
    from docx import Document
    for name,raw in outputs.items():
        if name.endswith('.docx'):
            paragraphs=[p.text for p in Document(io.BytesIO(raw)).paragraphs]
            group=next(p for p in paragraphs if '[214]马' in p and '[51]怕' in p)
            assert group.index('[214]')<group.index('[55]')<group.index('[51]')<group.index('[35]')
            assert '妈 巴 妈' in group and '麻文 鼻鼻化' in group
