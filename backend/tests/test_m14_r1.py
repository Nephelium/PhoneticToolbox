import io
import json
import unicodedata
import pytest
from docx import Document
from openpyxl import Workbook, load_workbook
from phonetic_core.transcription.phonology.parser import PhonologyInductionParser
from ptb_worker.m14_table import load_v2, inspect
from ptb_worker.m14_jobs import execute

OPTIONS=dict(computation_revision='m14/2',character_column=1,ipa_column=2,note_column=3,start_row=2,table_index=0,encoding='auto',delimiter='auto',consonant_only_as_zero_initial=True)

@pytest.mark.parametrize('ipa,zero,expected',[
 ('t͡s55',True,('Ø','t͡s','55')),('t͜sʰ55',False,('t͜sʰ','','55')),
 ('t͡ʃa35',True,('t͡ʃ','a','35')),('ã35',False,('Ø','ã','35')),
 ('tsã35',True,('ts','ã','35')),('ɹa55',True,('ɹ','a','55')),
 ('ɻa55',True,('ɻ','a','55')),('zɻ55',True,('z','ɻ','55')),
 ('ma⁵⁵',True,('m','a','55')),('m̩55',True,('Ø','m̩','55')),
 ('m̩55',False,('m̩','','55')),('kŋ̩55',True,('k','ŋ̩','55')),
 ('ˈma55',True,('m','a','55')),('/ ˌtsã35 /',True,('ts','ã','35')),
 ('cç55',True,('Ø','cç','55')),('cç55',False,('cç','','55')),
])
def test_r1_parse_independent_expected(ipa,zero,expected):
    value=PhonologyInductionParser('m14/2').parse(ipa,zero)
    assert (value.initial,value.final,value.tone)==expected

def test_r1_preserves_legacy_and_raw_ipa():
    assert PhonologyInductionParser().parse('ã35',False).initial=='ã'
    rows,_=load_v2('字,IPA,备注\n妈,tsã35,NA'.encode(),'a.csv',OPTIONS)
    assert rows[0].ipa=='tsã35' and rows[0].note=='NA'
    parsed=PhonologyInductionParser('m14/2')
    assert parsed.parse('tsã35')==parsed.parse('tsã35')

def test_r1_empty_columns_quotes_and_multi_syllable_diagnostics():
    raw='字\tIPA\t备注\n\tma55\t备注\n妈\tma55\t\n麻\tma35\t"逗号,及换行\n备注"\n多\tma55 pa21\t'.encode()
    rows,d=load_v2(raw,'a.tsv',OPTIONS)
    assert [r.character for r in rows]==['妈','麻']
    assert rows[1].note=='逗号,及换行\n备注'
    assert d['source_rows']==[3,4] and [r['row'] for r in d['skipped']]==[1,2,6]
    with pytest.raises(ValueError,match='start_inside_record'):load_v2(raw,'a.tsv',dict(OPTIONS,start_row=5))

@pytest.mark.parametrize('encoding',['utf-8-sig','gb18030','utf-16'])
def test_r1_encoding_mapping_start_row(encoding):
    raw='说明\r\n编号;备注;音标;字头\r\n1;甲;ma55;妈\r\n2;;ma35;麻'.encode(encoding)
    rows,d=load_v2(raw,'a.txt',dict(OPTIONS,encoding=encoding,delimiter='semicolon',character_column=4,ipa_column=3,note_column=2,start_row=3))
    assert [(r.character,r.ipa,r.note) for r in rows]==[('妈','ma55','甲'),('麻','ma35','')]
    assert d['source_rows']==[3,4]

@pytest.mark.parametrize('suffix',['xlsx','docx'])
def test_r1_select_second_table_and_sample(suffix):
    buf=io.BytesIO()
    if suffix=='xlsx':
        w=Workbook();w.active.append(['不选','na0']);sheet=w.create_sheet('调查');sheet.append(['备注','字','IPA']);sheet.append(['甲','妈','ma55']);w.save(buf)
    else:
        w=Document();w.add_table(rows=1,cols=2).cell(0,0).text='不选';sheet=w.add_table(rows=2,cols=3)
        for r,values in enumerate([['备注','字','IPA'],['甲','妈','ma55']]):
            for c,value in enumerate(values):sheet.cell(r,c).text=value
        w.save(buf)
    opts=dict(OPTIONS,table_index=1,character_column=2,ipa_column=3,note_column=1)
    sample=inspect(buf.getvalue(),'a.'+suffix,opts)
    assert len(sample['tables'])==2 and sample['sample'][1]['cells']==['甲','妈','ma55']
    rows,_=load_v2(buf.getvalue(),'a.'+suffix,opts)
    assert (rows[0].character,rows[0].ipa,rows[0].note)==('妈','ma55','甲')

def test_r1_contract_bad_columns():
    from ptb_api.m14_models import M14Config
    with pytest.raises(ValueError):M14Config(action='preview',**dict(OPTIONS,character_column=2))
    assert M14Config(action='inspect',**dict(OPTIONS,character_column=2)).action=='inspect'

def test_r1_import_bad_columns():
    with pytest.raises(ValueError,match='column_selection'):load_v2('妈,ma55'.encode(),'a.csv',dict(OPTIONS,start_row=1,ipa_column=1))

def test_r1_missing_ipa_never_uses_note_and_tone_only_is_diagnosed():
    raw='字,IPA,备注\n妈,,ma55\n调,⁵⁵,甲\n重,ˈ,乙\n麻,ma35,丙'.encode()
    rows,d=load_v2(raw,'a.csv',OPTIONS)
    assert [r.character for r in rows]==['麻'] and d['source_rows']==[5]
    assert [r['row'] for r in d['skipped']]==[1,2,3,4]

def test_r1_valid_row_limit_without_header():
    raw=('妈\tma55\t\n'*10001).encode()
    with pytest.raises(ValueError,match='m14_row_budget'):load_v2(raw,'a.tsv',dict(OPTIONS,start_row=1))

def test_r1_actual_export_matches_all_records_and_tone_order():
    raw='字,IPA,备注\n妈,ma55,甲\n麻,ma35,乙\n妈,ma55,甲\n鼻,tsã35,丙'.encode()
    preview=json.loads(execute(raw,'a.csv',dict(OPTIONS,action='preview'))['m14-preview.json'])
    assert preview['computation_revision']=='m14/2' and preview['diagnostics']['source_rows']==[2,3,4,5]
    cfg=preview['config'];cfg['tone_order']=['35','55'];cfg['tone_map']={'35':'阳','55':'阴'}
    result=execute(raw,'a.csv',dict(OPTIONS,action='export',settings=cfg,font=dict(schema_version='font/1',zh='宋体',latin='Times New Roman',ipa='Doulos SIL',size_px=14)))
    assert len(result)==3
    x=load_workbook(io.BytesIO(result['同音字表_二维表.xlsx']),rich_text=True).active
    m=next(c.column for c in x[1] if c.value=='m');a=next(r for r in range(2,x.max_row+1) if x.cell(r,1).value=='a')
    assert str(x.cell(a,m).value)=='[阳]麻乙 [阴]妈甲妈甲'
    for name,value in result.items():
        if name.endswith('.docx'):
            paragraphs=[p.text for p in Document(io.BytesIO(value)).paragraphs]
            assert any('[阳]麻乙[阴]妈甲 妈甲' in p for p in paragraphs)
