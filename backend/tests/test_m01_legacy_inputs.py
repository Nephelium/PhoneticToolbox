"""v2 writer-shaped inputs, malformed files and exact original-frame cuts."""
import io
import pickle
import sqlite3
from dataclasses import replace
import numpy as np
import pytest
from ptb_worker.io.limits import Limits, FormatError, LimitError
from ptb_worker.io.lip import convert_local_legacy_lip, decode_lip
from ptb_worker.io.legacy_parameters import read_legacy_parameters


@pytest.mark.parametrize('protocol',[2,3,4,5])
def test_v2_numpy_lip_and_companion(protocol):
    data={'absolute_timestamps':list(np.array([100.,100.25,100.5])),
          'relative_times':np.array([0.,.25,.5]),'open':[1.,np.nan,np.inf],
          'landmarks':[np.ones((478,3),dtype=np.float32)],
          'metadata':{'audio_first_frame_time':None,'lip_manual_offset':np.float64(.125)}}
    raw=pickle.dumps(data,protocol=protocol)
    result=decode_lip(convert_local_legacy_lip(raw,companion=pickle.dumps({'start_time':99.})))
    assert result['metadata']=={'audio_first_frame_time':99.,'lip_manual_offset':.125}
    assert result['relative_times']==[0.,.25,.5]
    assert result['absolute_timestamps']==[100.,100.25,100.5]
    assert np.isnan(result['open'][1]) and np.isposinf(result['open'][2])
    assert 'landmarks' not in result and raw==pickle.dumps(data,protocol=protocol)
    if protocol==2:
        # NumPy 1 used the public core module spelling; protocol 2 GLOBAL has
        # no frame/string length fields to rewrite when checking that alias.
        earlier=raw.replace(b'numpy._core.',b'numpy.core.')
        assert decode_lip(convert_local_legacy_lip(earlier))['relative_times']==[0.,.25,.5]


def test_numpy_object_constructor_cycle_and_expansion_rejected(tmp_path):
    marker=tmp_path/'must-not-exist'
    class Attack:
        def __reduce__(self):return (eval,(f"open({str(marker)!r},'w').write('bad')",))
    cycle=[];cycle.append(cycle);shared=[0.,.1]
    for value in (Attack(),{'relative_times':np.array([Attack()],dtype=object)},
                  {'relative_times':cycle},{'relative_times':shared,'open':shared},{'relative_times':np.array([1+2j])}):
        with pytest.raises((FormatError,LimitError)):convert_local_legacy_lip(pickle.dumps(value))
    assert not marker.exists()
    with pytest.raises((FormatError,LimitError)):
        convert_local_legacy_lip(pickle.dumps({'relative_times':np.arange(20.)}),replace(Limits(),samples=10))
    with pytest.raises(FormatError):convert_local_legacy_lip(pickle.dumps({'relative_times':[0.]})+b'x')


def legacy_pair():
    from openpyxl import Workbook
    columns=['Time_s','pF0','textgrid_音节'];rows=[[0.,120.,'ɑ̃˥'],[.1,None,''],[.5,'inf','=1+1'],[.9,'-inf','β']]
    wb=Workbook();ws=wb.active;ws.append(columns)
    for row in rows:
        ws.append(row)
        for cell in ws[ws.max_row]:
            if isinstance(cell.value,str):cell.data_type='s'
    out=io.BytesIO();wb.save(out);wb.close()
    db=sqlite3.connect(':memory:')
    try:
        db.execute('CREATE TABLE params (Time_s REAL,pF0 REAL,"textgrid_音节" TEXT)')
        db.executemany('INSERT INTO params VALUES (?,?,?)',[(t,float(v) if v else v,l) for t,v,l in rows]);db.commit()
        return out.getvalue(),db.serialize()
    finally:db.close()


@pytest.mark.parametrize('index,name',[(0,'声调.xlsx'),(1,'声调.ptb.sqlite')])
def test_old_tables_preserve_columns_frames_nonfinite_and_literal_labels(index,name):
    value=read_legacy_parameters(legacy_pair()[index],name)
    assert value['columns']==['Time_s','pF0','textgrid_音节']
    assert value['kinds']==['number','number','text']
    assert value['rows'][2]==[.5,'+Infinity','=1+1']
    assert value['rows'][3]==[.9,'-Infinity','β']


def test_tables_reject_formulas_unsorted_times_views_and_budget():
    from openpyxl import Workbook
    for rows in ([['Time_s','pF0'],[0.,'=1+1']], [['Time_s','pF0'],[1.,1.],[0.,2.]], [['Time_s','Time_s'],[0.,0.]]):
        wb=Workbook();ws=wb.active
        for row in rows:ws.append(row)
        stream=io.BytesIO();wb.save(stream);wb.close()
        with pytest.raises(FormatError):read_legacy_parameters(stream.getvalue(),'a.xlsx')
    db=sqlite3.connect(':memory:')
    try:
        db.execute('CREATE VIEW params AS SELECT 0 AS Time_s');db.commit()
        with pytest.raises(FormatError):read_legacy_parameters(db.serialize(),'a.ptb.sqlite')
    finally:db.close()
    for raw,name in zip(legacy_pair(),['a.xlsx','a.ptb.sqlite']):
        with pytest.raises(LimitError):read_legacy_parameters(raw,name,replace(Limits(),cells=5))
    with pytest.raises(LimitError):read_legacy_parameters(legacy_pair()[0],'a.xlsx',replace(Limits(),xml_bytes=10))


@pytest.mark.parametrize('index,name',[(0,'声调.xlsx'),(1,'声调.ptb.sqlite')])
def test_real_child_historical_cuts_preserve_source_and_samples(tmp_path,index,name):
    import hashlib
    from test_m01_segments import audio,grid,dead
    from ptb_worker.segmentation import prepare_segments
    from ptb_worker.io.scratch import Scratch
    from scipy.io import wavfile
    wav,samples=audio();raw=legacy_pair()[index];started=[]
    with Scratch(tmp_path,10_000_000) as scratch:
        bundle=prepare_segments(wav,grid(),'音节',scratch,legacy_result=raw,legacy_name=name,on_started=started.append)
        assert scratch.used==0
    assert started and dead(started[0])
    assert bundle.manifest['parent_result_sha256'] is None
    assert bundle.manifest['legacy_result']=={'name':name,'sha256':hashlib.sha256(raw).hexdigest(),'provenance':'user_associated_unverified'}
    np.testing.assert_array_equal(wavfile.read(io.BytesIO(bundle.payloads[3]))[1],samples[500:])
    table=read_legacy_parameters(bundle.payloads[5],'cut.ptb.sqlite')
    assert table['columns']==['Time_s','Source_Time_s','pF0','textgrid_音节']
    assert table['rows']==[[0.,.5,'+Infinity','=1+1'],[.4,.9,'-Infinity','β']]


def test_legacy_source_contract_cannot_impersonate_parent_or_enter_analysis():
    from pydantic import ValidationError
    from ptb_api.acoustic_batch_models import BatchRequest
    ids=[{'asset_id':f'00000000-0000-4000-8000-00000000000{i}','sha256':str(i)*64} for i in range(1,5)]
    body=dict(project_id=ids[0]['asset_id'],operation='textgrid_segment',idempotency_key='legacy-guard',layer='音节',inputs=[dict(audio=ids[0],textgrid=ids[1],legacy_result=ids[2])])
    assert BatchRequest(**body).inputs[0].legacy_result
    body['inputs'][0]['parent_result']=ids[3]
    with pytest.raises(ValidationError):BatchRequest(**body)
    del body['inputs'][0]['parent_result'];body.update(operation='acoustic_analysis',layer=None,config={})
    with pytest.raises(ValidationError):BatchRequest(**body)


def test_conversion_actual_owned_process():
    import struct
    from ptb_worker.legacy_conversion import convert
    raw=pickle.dumps({'relative_times':np.array([0.,.1]),'open':np.array([1.,2.])})
    assert decode_lip(convert(struct.pack('<II',len(raw),0)+raw))['open']==[1.,2.]


@pytest.mark.parametrize('case',['corrupt','formula','time','segmented','budget'])
def test_real_child_legacy_errors_return_no_partial_bundle(tmp_path,case):
    from openpyxl import Workbook
    from test_m01_segments import audio,grid
    from ptb_worker.segmentation import prepare_segments,SEGMENT_LIMITS
    from ptb_worker.io.scratch import Scratch
    wav,_=audio();wb=Workbook();sheet=wb.active
    sheet.append(['Time_s','Source_Time_s' if case=='segmented' else 'pF0'])
    sheet.append([2. if case=='time' else 0.,'=1+1' if case=='formula' else 120.])
    stream=io.BytesIO();wb.save(stream);wb.close()
    raw=b'bad workbook' if case=='corrupt' else stream.getvalue()
    expected='legacy_parameter_time_mismatch' if case in ('time','segmented') else 'legacy_parameter_budget' if case=='budget' else 'legacy_parameter_invalid'
    with Scratch(tmp_path,10_000_000) as scratch:
        with pytest.raises(FormatError,match=expected):
            prepare_segments(wav,grid(),'音节',scratch,legacy_result=raw,legacy_name='old.xlsx',limits=replace(SEGMENT_LIMITS,cells=2) if case=='budget' else SEGMENT_LIMITS)
        assert scratch.used==0


def test_xml_extent_entities_and_sqlite_generated_columns_are_rejected():
    from zipfile import ZipFile,ZIP_DEFLATED
    original=legacy_pair()[0]
    for xml in (b'<!DOCTYPE worksheet [<!ENTITY x "abc">]><worksheet/>',
                b'<Relationships><Relationship TargetMode="External" Target="https://invalid.test/"/></Relationships>',
                b'<worksheet><sheetData><row r="9999999"/></sheetData></worksheet>',
                b'<worksheet><sheetData><row r="1"><c r="ZZZ1"/></row></sheetData></worksheet>',
                '<!DOCTYPE worksheet [<!ENTITY x "abc">]><worksheet/>'.encode('utf-16')):
        out=io.BytesIO()
        with ZipFile(io.BytesIO(original)) as source,ZipFile(out,'w',ZIP_DEFLATED) as target:
            for entry in source.infolist():target.writestr(entry.filename,xml if entry.filename=='xl/worksheets/sheet1.xml' else source.read(entry))
        with pytest.raises((FormatError,LimitError)):read_legacy_parameters(out.getvalue(),'a.xlsx')
    db=sqlite3.connect(':memory:')
    try:
        db.execute('CREATE TABLE params(Time_s REAL,pF0 REAL AS (Time_s*10))');db.execute('INSERT INTO params(Time_s) VALUES(0)');db.commit()
        with pytest.raises(FormatError):read_legacy_parameters(db.serialize(),'a.ptb.sqlite')
    finally:db.close()
